"""
DNL Automation Tool — Single-file backend.
Consolidates: app.py, app_legacy.py, pipeline.py, google_news.py,
              pr_scraper.py, ai_filter.py, filters.py, exporter.py,
              config.py, logger.py

Run with:
    streamlit run main.py

Environment (.env):
    GROQ_API_KEY      — Groq API key for AI relevancy filter
    GROQ_MODEL        — optional model override (default: llama-3.3-70b-versatile)
    DNL_LOGIN_EMAIL   — login email   (default: dnladmin@automation.in)
    DNL_LOGIN_PASSWORD— login password (default: 123456789)
"""

# ============================================================
# STDLIB / THIRD-PARTY IMPORTS
# ============================================================
import json
import logging
import os
import re
import runpy
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple, Union
from urllib.parse import quote_plus, urljoin, urlparse

import feedparser
import pandas as pd
import requests
import streamlit as st
import streamlit.components.v1 as components
from bs4 import BeautifulSoup
from dateutil import parser as dateparser
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

load_dotenv()

# ============================================================
# LOGGER
# ============================================================
LOG_DIR = "logs"
os.makedirs(LOG_DIR, exist_ok=True)
_log_path = os.path.join(LOG_DIR, f"dnl_{datetime.now().strftime('%Y%m%d')}.log")

logger = logging.getLogger("dnl")
logger.setLevel(logging.INFO)

if not logger.handlers:
    _formatter = logging.Formatter(
        "%(asctime)s | %(levelname)-7s | %(name)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    _file_h = logging.FileHandler(_log_path, encoding="utf-8")
    _file_h.setFormatter(_formatter)
    _file_h.setLevel(logging.INFO)
    logger.addHandler(_file_h)

    _console_h = logging.StreamHandler()
    _console_h.setFormatter(_formatter)
    _console_h.setLevel(logging.INFO)
    logger.addHandler(_console_h)
    logger.propagate = False


# ============================================================
# CONFIG  (edit these lists to add companies / products / PR sites)
# ============================================================

COMPANIES = [
    "A Menarini Diagnostics","Abbott","AC Immune","Aesku Diagnostics",
    "aimmune therapeutics","Alcor Scientific","Alerje","Align ENT + Allergy",
    "AliveDX","ALK","Alladapt Immunotherapeutics","Allegra","AllerGenis",
    "Allergy Amulet","Allergy Partners","Allergy therapeutics","Allermi",
    "AllerSmart","Allervie Health","Aravax","Asit biotech","ATANIS Biotech AG",
    "Augurex Life Sciences","Autobio Diagnostics","Autu Bio",
    "Axim Biotechnologies","Bayer Healthcare LLC","BBI Solutions",
    "BD Biosciences","Beckman Coulter","Bio-Rad","Biosynex",
    "Bio-Techne Corporation","Bredis Healthcare","BSBE Biotech","Buhlmann",
    "BÜHLMANN Laboratories","C2N Diagnostics","Camallergy",
    "Canon Medical Systems","CellaVision","Celltrion","Cerascreen",
    "CF PharmTech","Circular Genomics","Cizzle Biotechnology",
    "C-Luminary Biotech","CND Life Sciences","Co-op","Creative Proteomics",
    "Crown Bioscience","CVS Health","Cytox","Danaher Corporation",
    "DBV technologies","Desentum","Diadem","DiagnaMed","DiamiR","DiaSorin",
    "Differentia Biotech","DOTS Technology Corp","Dr. Fooke","Droege Group",
    "Durin Technologies","Eiken Chemical","EKF Diagnostics","empowerDx",
    "Enterome","Euro Diagnostica","Euroimmun","Eurospital","Evotec SE",
    "Exagen","Excellergy","FAES Farma S.A","Flappd Tech","FSA","Fujirebio",
    "Gentian Diagnostics","Global Medical Technologies",
    "Gold Standard Diagnostics","Grifols",
    "Guangzhou Leide Bioscience Co., Ltd","Henley Business School",
    "Hitachi Chemical Diagnostics","Hitachi Chemical Diagnostics Systems",
    "HOB Biotech","Hongrui Taijie","Hotgen Bio","Human Diagnostics","Hycor",
    "HYCOR Biomedical","Iason","IgGenix","Imcyse","Immuno Concepts",
    "Immuno Diagnostic System","InBio","Infinity BiologiX","Inimmune",
    "Inmune Bio","Inova Diagnostics","Intrommune Therapeutics","Invitros",
    "JIUHUA","JSR Life Sciences","Jubilant Pharmova","Kenota Health",
    "KingMed Diagnostics","Koneksa","LG Chem","MADX","MannKind","MBL",
    "MedicalSystem Biotechnology","Mediwiss Analytics GmbH",
    "Medsource Ozone Biomedicals","Medtronic","Menarini","Merck",
    "Meso Scale Diagnostics","Metabolon","MicroTech Medical","Mikrogen",
    "Minaris Medical","Mindray","Mitsubishi Gas Chemical","Morepen",
    "MYND Diagnostics","NA","NADMED","Nanchang Baiji Medicine Technology",
    "Nanjing Vazyme Biotech","Nectar","Neogen","Neosinus Health",
    "Nexkin Medical","Nippon Chemiphar","ObvioHealth","Octave Bioscience",
    "Omega Diagnostics","OMRF","Oncimmune","Oncologica","Optina Diagnostics",
    "Orgentec","Ortho Clinical Diagnostics","Osler","Oumeng",
    "PEPperPRINT GmbH","Perspectum","Potens Allergy","PrecisionLife",
    "Progentec Diagnostics, Inc.","Prolight Diagnostics",
    "Prometheus Laboratories","Quibim","Quidel","QuidelOrtho","Quotient",
    "R-Biopharm","Reacta Biotech","Resonac","Revelation Biosciences",
    "Revvity","Roche Diagnostics","RSR","Sebia","Sharay","Shimadzu Corp.",
    "Showa Denko Materials","Siemens Healthineers","Snibe",
    "Stallergenes Greer","Sugentech","Svar Lifesciences","Tellgen",
    "Theradiag","Trinity Biotech","Ukko","VitalFlo","Werfen","WM Partners",
    "Wondfo Bio","WuXi Diagnostics","YHLO","ZEUS Scientific",
]

PRODUCT_CATALOG = {
    "A. Menarini Diagnostics": [
        "ZENIT","Zenit PRO","Zenit fast","Zenit Reagents","Zenit HUB",
        "Zenit Prime","CHORUS","CHORUS TRIO",
    ],
    "Abbott": [
        "Alinity i","ARCHITECT","ARCHITECT i1000SR","ARCHITECT i2000SR",
        "Alinity ci-series","CHEMIFLEX","Lab Central",
    ],
    "AESKU.GROUP": [
        "AESKULISA","AESKUSLIDES","AESKUBLOTS","AESQC","HELIOS",
        "AKLIDES","akiron","DST","ATLAS",
    ],
    "AliveDx": ["MosaiQ","LumiQ","Alba"],
    "Autobio": [
        "AutoLumo","AutoLumo A1000","AutoLumo A1860","AutoLumo A2000",
        "AutoLumo A2000 Plus","AutoLumo A6200","AutoLumo S900",
        "Autobio CLIA","CLIA Microparticles",
    ],
    "Beckman Coulter": [
        "Access","Access 2","DxI","DxI 9000","DxI 9000 Access",
        "UniCel DxI","UniCel DxC","DxC 500i",
    ],
    "Bio-Rad": [
        "BioPlex 2200","BioPlex 2200 System","BioPlex 2200 Autoimmune Panels",
        "Kallestad","EVOLIS","Platelia","BioPlex",
    ],
    "BUHLMANN Laboratories": [
        "CAST","CAST ELISA","Flow CAST","GanglioCombi","Quantum Blue",
        "BÜHLMANN fCAL","fCAL turbo","IBDoc","Calex","CalApp",
    ],
    "C-Luminary Biotech / Sharay": [
        "Sharay","Sharay 1600","Sharay 6000","Sharay 8000","Sharay 800",
        "C-Luminary CLIA","Xieguang Bio",
    ],
    "DiaSorin": [
        "LIAISON","LIAISON XL","LIAISON XS","LIAISON IQ","xMAP",
        "xMAP INTELLIFLEX","Simplexa","Luminex","VERIGENE",
    ],
    "Dr. Fooke": [
        "ALLERG-O-LIQ","ALLERG-O-SYSTEM","ALFA","ALFA Reader",
        "ALLERG-O-WIN","ARTHUS","LFA Reader","Dr. Fooke Laboratories",
        "Fooke Labs",
    ],
    "Eurospital": [
        "Eu-tTG","α-GliaPep","alpha-GliaPep","α-Gliatest","alpha-Gliatest",
        "Xeliac","Xeliac Test Pro","Xeliac Test Professional","XeliGen","EAL",
        "ECD","Eurospital Autoimmunity Product Line","Easy-NAT",
        "Eu-King's STAR","Calprest",
    ],
    "Gold Standard Diagnostics": [
        "Hurricane Dx","AESKULISA","NovaLisa","The Bolt","ThunderBolt",
        "AIX1000","NovaTec","Virotech","Gold Standard Diagnostics Frankfurt",
    ],
    "HOB Biotech": [
        "BioCLIA","BioCLIA 500","BioCLIA 6500","BioCLIA Allergy",
        "BioCLIA Autoimmune","BioLINE","BioCLIA 1200","REAST","HOB Allergy",
        "HOB Autoimmune","EUROHOB",
    ],
    "HYCOR": ["NOVEOS","HYTEC","HYTEC 288 Plus","NOVEOS Allergens"],
    "Invitros": [
        "AdvanSure AlloScreen","AdvanSure AlloScreen Max108 Panel","AdvanSure",
        "AdvanSure i3 TB-IGRA","AdvanSure RV-Plus","AdvanSure E3 System",
    ],
    "MADx": [
        "ALEX","ALEX²","ALEX2","ALEX³","ALEX3","Allergy Xplorer","FOX",
        "Food Xplorer","ImageXplorer","Image Xplorer","MAX 9k","MAX45k",
        "MAX 45k","Raptor Server","MADx Cloud","MacroArray Diagnostics",
    ],
    "MBL Life Science": [
        "MESACUP","MESACUP Series","MEBLux","STACIA","MBL IVD",
        "MBL Autoimmune",
    ],
    "Nippon Chemiphar": [
        "DropScreen","DropScreen ST-1","Drop Screen","IgE NC","R-NanoBio",
    ],
    "QuidelOrtho": [
        "MICROVUE","MICROVUE Complement Multiplex","THYRETAIN","VITROS",
        "ELISA Kits","TECOMedical","QuidelOrtho Specialty Products Group","SPG",
    ],
    "Revvity / EUROIMMUN / IDS": [
        "EUROLINE","EUROBlotOne","EUROLineScan","EUROPattern","IDS-iSYS",
        "IDS-i10","IDS i20","IDS-i20","EUROLabWorkstation","EUROLabOffice 4.0",
        "EUROIMMUN Analyzer I","IF Sprinter","Sprinter XL",
    ],
    "Roche": [
        "Elecsys","cobas e","cobas e 411","cobas e 402","cobas e 801",
        "cobas pure","cobas pro","ECLIA",
    ],
    "Sebia": [
        "Alegria","Alegria 2","Alegria Monotest","SMC Technology","ORGENTEC",
        "myAlegria","ELISA","Immunofluorescence","Immunoblot",
    ],
    "Siemens Healthineers": [
        "IMMULITE","IMMULITE 2000","IMMULITE 2000 XPi","IMMULITE 1000",
        "3gAllergy","IMMULITE 3gAllergy","Atellica","Atellica Solution",
        "Atellica CI","ADVIA Centaur","Dimension Vista","Dimension EXL",
    ],
    "Snibe Diagnostics": [
        "MAGLUMI","MAGLUMI X3","MAGLUMI X6","MAGLUMI X8","MAGLUMI 600",
        "MAGLUMI 800","MAGLUMI 2000","MAGLUMI 2000 Plus","MAGLUMI M Series",
        "Biolumi CX8","Biolumi CX Solution","X-TECH",
    ],
    "Sugentech": [
        "SGTi-Allergy Screen","SGTi-Allergy Screen PLUS",
        "SGTi-Allergy PLUS 60","S-Blot 3","S-Blot 3 PLUS","S-Blot 2",
        "S-Blot 2 Easy","S-Blot 2 Easy PLUS","S-Blot PLUS","INCLIX F-100",
        "INCLIX",
    ],
    "Theradiag": [
        "FIDIS","Lisa Tracker","LISA TRACKER","i-Track10","i-Tracker","xMAP",
        "Theranostics","Biosynex","Luminex",
    ],
    "Trinity Biotech": [
        "ImmuGlo","Mardx","ImmuLisa","Captia","ImmuBlot","Marstripe",
        "ImmcoStripe","Amerlex","Immco","Immco Diagnostics",
        "Trinity Autoimmune",
    ],
    "Werfen": [
        "Aptiva","BIO-FLASH","QUANTA Flash","QUANTA Lite","NOVA View",
        "NOVA Lite","QUANTA-Lyser 3000","QUANTA Link","Integrated Lab+",
        "Inova Diagnostics",
    ],
    "YHLO": [
        "iFlash","iFlash 1200","iFlash 1800","iFlash 3000","iFlash 9000",
        "iStar 500","CLIA Solution","LIA Solution",
    ],
}

PRODUCTS = [
    {"company": company, "name": product}
    for company, products in PRODUCT_CATALOG.items()
    for product in products
]

PR_WEBSITES = [
    {"company": "Aesku Diagnostics",      "url": "https://www.aesku.com/"},
    {"company": "Abbott",                  "url": "https://abbott.mediaroom.com/press-releases"},
    {"company": "Danaher",                 "url": "https://www.danaher.com/newsroom"},
    {"company": "Beckman Coulter",         "url": "https://www.beckmancoulter.com/about-beckman-coulter/newsroom/press-releases"},
    {"company": "Bio-Rad (IR)",            "url": "https://investors.bio-rad.com/press-releases/default.aspx"},
    {"company": "Diasorin",                "url": "https://int.diasorin.com/en/investors/finance/press-releases"},
    {"company": "Dr. Fooke",               "url": "https://www.fooke-labs.de/veranstaltungen?lang=en"},
    {"company": "SVAR Lifesciences",       "url": "https://www.svarlifescience.com/news"},
    {"company": "Euroimmun",               "url": "https://www.euroimmun.us/recent-news"},
    {"company": "Eurospital",              "url": "https://www.eurospital.com/news/"},
    {"company": "HOB Biotech",             "url": "http://en.hob-biotech.com/media-55.html"},
    {"company": "Hycor Biomedical",        "url": "https://www.hycorbiomedical.com/news/categories/hycor-news-archive-2023"},
    {"company": "MBL",                     "url": "https://www.mblbio.com/e/ir/press.html#"},
    {"company": "JSR Corporation",         "url": "https://www.jsr.co.jp/jsr_e/news/2026/"},
    {"company": "QuidelOrtho",             "url": "https://ir.quidelortho.com/home/default.aspx"},
    {"company": "C-Luminary",              "url": "https://www.c-luminary.com/list-46-1.html"},
    {"company": "Nippon Chemiphar",        "url": "https://www.chemiphar.co.jp/english/"},
    {"company": "Human Diagnostics",       "url": "https://www.human.de/about-human/overview-human/news"},
    {"company": "Autobio Diagnostics",     "url": "https://en.autobio.com.cn/News/index/fid/3/cid/2/id/142.html"},
    {"company": "Siemens Healthineers",    "url": "https://www.siemens-healthineers.com/press/releases"},
    {"company": "Menarini",                "url": "https://www.menarini.com/en-us/news"},
    {"company": "Roche",                   "url": "https://www.roche.com/media/releases"},
    {"company": "Sebia",                   "url": "https://www.sebia.com/en-in/resources/"},
    {"company": "Revvity",                 "url": "https://news.revvity.com/press-announcements/press-releases/default.aspx"},
    {"company": "Zeus Scientific",         "url": "https://www.zeusscientific.com/news-events"},
    {"company": "Beckman (India)",         "url": "https://www.mybeckman.in/news"},
    {"company": "AliveDx",                 "url": "https://alivedx.com/news/"},
    {"company": "Werfen",                  "url": "https://www.werfen.com/en/news"},
    {"company": "Theradiag",               "url": "https://www.theradiag.com/en/category/press-releases/"},
    {"company": "Trinity Biotech",         "url": "https://www.trinitybiotech.com/category/press-releases/"},
    {"company": "BUHLMANN Laboratories",   "url": "https://www.buhlmannlabs.ch/news/"},
    {"company": "IDS plc",                 "url": "https://www.idsplc.com/blog/"},
    {"company": "Gold Standard Diagnostics","url": "https://www.goldstandarddiagnostics.com/news-event/category/news"},
    {"company": "R-Biopharm",              "url": "https://clinical.r-biopharm.com/news/"},
    {"company": "MADx",                    "url": "https://www.macroarraydx.com/news/overview"},
    {"company": "Aimmune",                 "url": "https://www.aimmune.com/news-releases"},
    {"company": "Stallergenes Greer",      "url": "https://www.stallergenesgreer.com/press-releases#"},
    {"company": "Allergy Therapeutics",    "url": "https://www.allergytherapeutics.com/investors/regulatory-news/"},
    {"company": "ALK-Abelló",             "url": "https://ir.alk.net/news-events/company-releases"},
    {"company": "DBV Technologies",        "url": "https://dbv-technologies.com/investor-overview/news/"},
    {"company": "Eurospital (EGL)",        "url": "https://egl.eurospital.com/en/news-and-events/"},
]

# Keywords loaded from keywords.txt if it exists, otherwise use defaults
def _load_keywords_from_file():
    kw_path = Path("keywords.txt")
    if not kw_path.exists():
        return None
    try:
        lines = kw_path.read_text(encoding="utf-8").splitlines()
        relevant, exclude, hot = [], [], []
        section = None
        for line in lines:
            line = line.strip()
            if not line or line.startswith("#"):
                if "RELEVANT" in line.upper():
                    section = "relevant"
                elif "EXCLUDE" in line.upper():
                    section = "exclude"
                elif "HOT" in line.upper():
                    section = "hot"
                continue
            if section == "relevant":
                relevant.append(line)
            elif section == "exclude":
                exclude.append(line)
            elif section == "hot":
                hot.append(line)
        return relevant or None, exclude or None, hot or None
    except Exception:
        return None

_kw_file = _load_keywords_from_file()

RELEVANT_KEYWORDS = (
    _kw_file[0] if _kw_file and _kw_file[0] else [
        "immunodiagnostic","allergy diagnostic","autoimmune diagnostic",
        "immunoassay","diagnostic platform","point-of-care","poc","ivd",
        "elisa","lateral flow","multiplex","serology",
        "allergy immunotherapy","allergen immunotherapy","scit","slit",
        "allergy therapeutic","allergy","autoimmune","immunology",
        "merger","acquisition","acquires","acquired","takeover","divest",
        "partnership","collaboration","agreement","alliance","joint venture",
        "product launch","new product","launches","launch","platform",
        "ce mark","ce-mark","fda clearance","fda approval","510(k)",
        "regulatory approval","regulatory","ema approval","market authorization",
        "conference","symposium","congress","eaaci","aaaai","acaai","escmid",
        "diagnostics","diagnostic","assay","reagent","analyzer",
        "laboratory","lab","clinical","test kit",
        "revenue","earnings","financial results",
        "ceo","cfo","appoints","appointment","leadership","executive",
    ]
)

EXCLUDE_KEYWORDS = (
    _kw_file[1] if _kw_file and _kw_file[1] else [
        "generic drug","biosimilar","chemotherapy","oncology drug",
        "cardiovascular drug","diabetes drug","obesity drug",
        "mental health","psychiatry","neurology drug","orthopedic",
        "dermatology drug","ophthalmology drug",
    ]
)

HOT_KEYWORDS = (
    _kw_file[2] if _kw_file and _kw_file[2] else [
        "merger","acquisition","acquires","acquired","takeover","divest",
        "fda approval","fda clearance","ce mark","510(k) clearance",
        "market authorization","regulatory approval",
        "strategic alliance","joint venture","major partnership",
        "platform launch","breakthrough","pivotal","landmark",
    ]
)

NEWS_TYPE_RULES = {
    "M&A": ["merger","acquisition","acquires","acquired","takeover","divest"],
    "Regulatory": [
        "regulatory","approval","cleared","clearance","authorized",
        "authorization","fda","ema","ce mark","ce-mark","510(k)",
    ],
    "Product Launch": [
        "launch","launches","new product","new platform","introduces",
        "unveiled","unveils",
    ],
    "Partnership": [
        "partnership","collaboration","agreement","alliance",
        "joint venture","co-develop","license","licensing",
    ],
    "Conference": [
        "conference","congress","symposium","expo",
        "eaaci","aaaai","acaai","ecp","esh","escmid",
    ],
    "Organizational": [
        "ceo","cfo","appoints","appointment","leadership",
        "executive","board","director","restructure",
    ],
    "Financial": [
        "revenue","earnings","results","financial","profit","growth","guidance",
    ],
}

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept": (
        "text/html,application/xhtml+xml,application/xml;q=0.9,"
        "image/avif,image/webp,*/*;q=0.8"
    ),
    "Accept-Language": "en-US,en;q=0.9",
}
REQUEST_TIMEOUT = 15


# ============================================================
# FILTERS
# ============================================================
def _contains_any(text: str, vocab) -> bool:
    t = (text or "").lower()
    return any(k.lower() in t for k in vocab)

def keyword_is_relevant(text: str) -> bool:
    if not text:
        return False
    if _contains_any(text, EXCLUDE_KEYWORDS):
        return False
    return _contains_any(text, RELEVANT_KEYWORDS)

def classify_news_type(text: str) -> str:
    t = (text or "").lower()
    for category, triggers in NEWS_TYPE_RULES.items():
        if any(trig.lower() in t for trig in triggers):
            return category
    return "Other"

def classify_hot(text: str) -> str:
    return "HOT" if _contains_any(text, HOT_KEYWORDS) else "Non-Hot"

def deduplicate(records: List[Dict]) -> List[Dict]:
    seen = set()
    out = []
    for r in records:
        entity = (r.get("screened_entity") or r.get("company") or "").lower().strip()
        category = (r.get("screening_category") or "").lower().strip()
        headline = (r.get("headline") or "").lower().strip()[:80]
        key = (category, entity, headline)
        if key in seen:
            continue
        seen.add(key)
        out.append(r)
    return out


# ============================================================
# GOOGLE NEWS
# ============================================================
GNEWS_TEMPLATE = "https://news.google.com/rss/search?q={query}&hl=en-US&gl=US&ceid=US:en"

def _build_query(company: str) -> str:
    raw = f'"{company}" diagnostics OR allergy OR immunology'
    return GNEWS_TEMPLATE.format(query=quote_plus(raw))

def _build_product_query(product: str) -> str:
    raw = f'"{product}" diagnostics OR allergy OR immunology'
    return GNEWS_TEMPLATE.format(query=quote_plus(raw))

def _parse_date(entry) -> Optional[date]:
    for field in ("published_parsed", "updated_parsed"):
        tup = entry.get(field)
        if tup:
            try:
                return datetime(*tup[:6]).date()
            except Exception:
                pass
    for field in ("published", "updated"):
        raw = entry.get(field)
        if raw:
            try:
                return dateparser.parse(raw).date()
            except Exception:
                pass
    return None

def _fetch_entity_news(
    entities: List[Union[str, Dict[str, str]]],
    start_date: date,
    end_date: date,
    query_builder: Callable[[str], str],
    screening_category: str,
    progress_callback: Optional[Callable[[str, int, int], None]] = None,
) -> List[Dict]:
    records: List[Dict] = []
    total = len(entities)
    for idx, entity in enumerate(entities, start=1):
        if isinstance(entity, dict):
            entity_name = entity["name"]
            owner = entity.get("company") or entity_name
        else:
            entity_name = entity
            owner = entity
        if progress_callback:
            try:
                progress_callback(entity_name, idx, total)
            except Exception:
                pass
        url = query_builder(entity_name)
        try:
            feed = feedparser.parse(url)
        except Exception as e:
            logger.warning(f"[GNews:{screening_category}] {entity_name}: feed parse failed: {e}")
            continue
        if getattr(feed, "bozo", False) and not feed.entries:
            continue
        matched = 0
        for entry in feed.entries:
            pub = _parse_date(entry)
            if pub is None:
                continue
            if pub < start_date or pub > end_date:
                continue
            headline = (entry.get("title") or "").strip()
            link = (entry.get("link") or "").strip()
            if not headline or not link:
                continue
            records.append({
                "company": owner,
                "screened_entity": entity_name,
                "screening_category": screening_category,
                "headline": headline,
                "date": pub.strftime("%Y-%m-%d"),
                "url": link,
                "source_type": "Google News",
            })
            matched += 1
        logger.info(f"[GNews:{screening_category}] {entity_name}: {matched}/{len(feed.entries)} in range")
    return records

def fetch_google_news(companies, start_date, end_date, progress_callback=None):
    return _fetch_entity_news(
        companies, start_date, end_date,
        query_builder=_build_query,
        screening_category="Competitor Screening",
        progress_callback=progress_callback,
    )

def fetch_product_news(products, start_date, end_date, progress_callback=None):
    return _fetch_entity_news(
        products, start_date, end_date,
        query_builder=_build_product_query,
        screening_category="Product Screening",
        progress_callback=progress_callback,
    )


# ============================================================
# PR SCRAPER
# ============================================================
_DATE_PATTERNS = [
    re.compile(r"\b(\d{4}-\d{2}-\d{2})\b"),
    re.compile(r"\b((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{1,2},?\s+\d{4})\b", re.IGNORECASE),
    re.compile(r"\b(\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4})\b", re.IGNORECASE),
    re.compile(r"\b(\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4})\b"),
]
_DATE_CLASS_HINTS = ("date","published","pub-date","pubdate","timestamp","time","posted",)
_NOISE_RE = re.compile(
    r"^(home|news|press|releases?|menu|search|login|sign in|sign up|subscribe|"
    r"contact|cookie|privacy|terms|imprint|read more|learn more|more|back|next|"
    r"previous|all|view all|see all|share|download|english|deutsch|français|"
    r"about|careers|investors|products|solutions|support|events)$",
    re.IGNORECASE,
)

def _try_parse(raw: str) -> Optional[date]:
    if not raw:
        return None
    raw = raw.strip()
    if len(raw) < 6 or len(raw) > 60:
        return None
    try:
        d = dateparser.parse(raw, fuzzy=True).date()
    except Exception:
        return None
    if d.year < 2000 or d.year > date.today().year + 1:
        return None
    return d

def _date_from_attrs(el) -> Optional[date]:
    if not hasattr(el, "get"):
        return None
    for attr in ("datetime","data-date","data-published","data-time","content"):
        v = el.get(attr)
        if v:
            d = _try_parse(v)
            if d:
                return d
    return None

def _date_from_text(text: str) -> Optional[date]:
    if not text:
        return None
    for pat in _DATE_PATTERNS:
        m = pat.search(text)
        if m:
            d = _try_parse(m.group(1))
            if d:
                return d
    return None

def _date_from_element(el, max_text: int = 1500) -> Optional[date]:
    if el is None or not hasattr(el, "find_all"):
        return None
    d = _date_from_attrs(el)
    if d:
        return d
    for t in el.find_all("time", limit=6):
        d = _date_from_attrs(t)
        if d:
            return d
        d = _try_parse(t.get_text(" ", strip=True))
        if d:
            return d
    for cand in el.find_all(True, limit=40):
        cls = " ".join(cand.get("class") or []).lower()
        if not any(hint in cls for hint in _DATE_CLASS_HINTS):
            continue
        d = _date_from_attrs(cand)
        if d:
            return d
        d = _try_parse(cand.get_text(" ", strip=True))
        if d:
            return d
    return _date_from_text(el.get_text(" ", strip=True)[:max_text])

def _find_date_for_anchor(a) -> Optional[date]:
    d = _date_from_element(a, max_text=400)
    if d:
        return d
    parent = a.parent
    for _ in range(4):
        if parent is None or parent.name in ("body","html"):
            break
        d = _date_from_element(parent, max_text=1500)
        if d:
            return d
        parent = parent.parent
    for sib_fn in (a.find_previous_sibling, a.find_next_sibling):
        sib = sib_fn()
        for _ in range(3):
            if sib is None:
                break
            d = _date_from_element(sib, max_text=400)
            if d:
                return d
            sib = sib.find_previous_sibling() if sib_fn == a.find_previous_sibling else sib.find_next_sibling()
    return None

def _is_internal_link(href: str, base_netloc: str) -> bool:
    if not href:
        return False
    if href.startswith(("javascript:","mailto:","tel:","#")):
        return False
    if href.startswith("/"):
        return True
    parsed = urlparse(href)
    if not parsed.netloc:
        return True
    return base_netloc.split(".")[-2:] == parsed.netloc.split(".")[-2:]

def _collect_anchor_candidates(soup, base_url: str) -> List:
    base_netloc = urlparse(base_url).netloc
    out = []
    seen_texts = set()
    for a in soup.find_all("a", href=True, limit=2000):
        href = a["href"].strip()
        if not _is_internal_link(href, base_netloc):
            continue
        headline = None
        for tag in ("h1","h2","h3","h4","h5"):
            h = a.find(tag)
            if h:
                txt = h.get_text(" ", strip=True)
                if 15 <= len(txt) <= 500:
                    headline = txt
                    break
        if not headline:
            txt = a.get_text(" ", strip=True)
            if 15 <= len(txt) <= 500 and not _NOISE_RE.match(txt):
                headline = txt
        if not headline:
            continue
        key = headline.lower()[:80]
        if key in seen_texts:
            continue
        seen_texts.add(key)
        abs_url = urljoin(base_url, href)
        out.append((a, headline, abs_url))
    return out

def _classify_failure(exc: Exception) -> Tuple[str, str]:
    if isinstance(exc, requests.HTTPError) and exc.response is not None:
        code = exc.response.status_code
        if code == 404:
            return ("HTTP 404 — Page not found", "URL returned 404.")
        if code == 403:
            return ("HTTP 403 — Forbidden", "Site blocked the request.")
        if 500 <= code < 600:
            return (f"HTTP {code} — Server error", "Remote server error.")
        return (f"HTTP {code}", f"Unexpected HTTP status {code}.")
    if isinstance(exc, requests.ConnectionError):
        s = str(exc).lower()
        if any(k in s for k in ("nodename","name or service","getaddrinfo")):
            return ("Domain unreachable (DNS)", "DNS lookup failed.")
        if "ssl" in s or "certificate" in s:
            return ("SSL / certificate error", "Invalid SSL certificate.")
        return ("Connection error", str(exc)[:160])
    if isinstance(exc, requests.Timeout):
        return ("Timeout", f"No response within {REQUEST_TIMEOUT}s.")
    if isinstance(exc, requests.TooManyRedirects):
        return ("Too many redirects", "Site redirected in a loop.")
    return (type(exc).__name__, str(exc)[:160])

def scrape_pr_websites(
    pr_websites, start_date, end_date,
    progress_callback=None, failure_callback=None,
    include_undated=True, undated_cap=15,
) -> Tuple[List[Dict], List[Dict]]:
    records: List[Dict] = []
    failures: List[Dict] = []
    total = len(pr_websites)
    today_str = date.today().strftime("%Y-%m-%d")

    def _record_failure(company, url, category, detail):
        f = {"company": company, "url": url, "category": category, "error": detail}
        failures.append(f)
        logger.warning(f"[PR FAIL] {company} ({url}): {category} — {detail}")
        if failure_callback:
            try:
                failure_callback(f)
            except Exception:
                pass

    for idx, site in enumerate(pr_websites, start=1):
        company = site["company"]
        url = site["url"]
        if progress_callback:
            try:
                progress_callback(company, idx, total)
            except Exception:
                pass
        host = urlparse(url).netloc.lower()
        if "linkedin.com" in host:
            _record_failure(company, url, "Non-compliant source (blocked)", "LinkedIn blocks scraping.")
            continue
        resp = None
        last_exc = None
        for attempt in range(2):
            try:
                resp = requests.get(url, headers=HEADERS, timeout=REQUEST_TIMEOUT)
                resp.raise_for_status()
                last_exc = None
                break
            except (requests.ConnectionError, requests.Timeout) as e:
                last_exc = e
                if attempt == 0:
                    time.sleep(1.5)
                    continue
            except Exception as e:
                last_exc = e
                break
        if last_exc is not None or resp is None:
            cat, detail = _classify_failure(last_exc or Exception("Unknown error"))
            _record_failure(company, url, cat, detail)
            continue
        try:
            soup = BeautifulSoup(resp.text, "lxml")
        except Exception as e:
            _record_failure(company, url, "HTML parse error", str(e)[:160])
            continue
        candidates = _collect_anchor_candidates(soup, url)
        in_range = 0
        dated_count = 0
        undated_pool = []
        for a, headline, link in candidates:
            d = _find_date_for_anchor(a)
            if d is None:
                undated_pool.append((headline, link))
                continue
            dated_count += 1
            if d < start_date or d > end_date:
                continue
            records.append({
                "company": company,
                "screened_entity": company,
                "screening_category": "PR Site Screening",
                "headline": headline,
                "date": d.strftime("%Y-%m-%d"),
                "url": link,
                "source_type": "PR Website",
                "date_estimated": False,
            })
            in_range += 1
        used_undated = 0
        if include_undated and dated_count == 0 and undated_pool:
            for headline, link in undated_pool[:undated_cap]:
                records.append({
                    "company": company,
                    "screened_entity": company,
                    "screening_category": "PR Site Screening",
                    "headline": headline,
                    "date": today_str,
                    "url": link,
                    "source_type": "PR Website",
                    "date_estimated": True,
                })
            used_undated = min(len(undated_pool), undated_cap)
        logger.info(f"[PR] {company}: candidates={len(candidates)} dated={dated_count} in_range={in_range} undated_used={used_undated}")
    return records, failures


# ============================================================
# AI FILTER (Groq)
# ============================================================
GROQ_ENDPOINT = "https://api.groq.com/openai/v1/chat/completions"
DEFAULT_GROQ_MODEL = "llama-3.3-70b-versatile"
REQUEST_TIMEOUT_SEC = 90

AI_SYSTEM_PROMPT = """You are an intelligence analyst preparing a Daily Newsletter (DNL) for a specialty diagnostics business unit focused on immunodiagnostics, allergy, and autoimmune testing.

Your job is to classify headlines as RELEVANT or NOT RELEVANT for the Daily Newsletter (DNL).

IN SCOPE (relevant):
- Immunodiagnostics, allergy diagnostics, autoimmune diagnostics
- Immunoassays (ELISA, lateral flow, multiplex, serology)
- Diagnostic platforms / analyzers / reagents / test kits
- Infectious disease diagnostics, point-of-care (POC) diagnostics, IVD
- Laboratory solutions
- Allergen immunotherapy / allergy therapeutics (ALK, Stallergenes Greer, Aimmune, DBV, Allergy Therapeutics are in scope)
- Strategic events on monitored companies: M&A, partnerships, regulatory approvals (FDA, CE, 510k), product launches, conference participation, leadership changes, strategic financial results

OUT OF SCOPE (not relevant):
- Generic pharma news (cardiovascular, oncology, diabetes, obesity)
- Biosimilars or unrelated drug approvals
- General hospital / healthcare system news
- Unrelated clinical trials for drugs
- Mental health / psychiatry / neurology / orthopedic / dermatology / ophthalmology drugs
- Generic medical news with no link to diagnostics or allergy

Return ONLY a JSON array, no prose, no markdown fences. Each element:
{"headline": "<exact headline you were given>", "relevant": true|false, "reason": "<one short sentence>"}
"""

def _extract_json_array(text: str):
    text = re.sub(r"^```(?:json)?", "", text.strip(), flags=re.IGNORECASE)
    text = re.sub(r"```$", "", text.strip())
    start = text.find("[")
    end = text.rfind("]")
    if start == -1 or end == -1 or end <= start:
        raise ValueError("no JSON array found in response")
    return json.loads(text[start: end + 1])

def _call_groq(api_key, model, user_msg, system_prompt=AI_SYSTEM_PROMPT) -> str:
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
    payload = {
        "model": model,
        "temperature": 0.1,
        "max_tokens": 4096,
        "response_format": {"type": "json_object"},
        "messages": [
            {"role": "system", "content": system_prompt + '\n\nReturn the JSON array under a key called "results".'},
            {"role": "user", "content": user_msg},
        ],
    }
    r = requests.post(GROQ_ENDPOINT, headers=headers, json=payload, timeout=REQUEST_TIMEOUT_SEC)
    if r.status_code >= 400:
        raise requests.HTTPError(f"HTTP {r.status_code} from Groq: {r.text[:500]}", response=r)
    data = r.json()
    return data["choices"][0]["message"]["content"]

def _parse_verdicts(text: str):
    text = text.strip()
    try:
        obj = json.loads(text)
        if isinstance(obj, list):
            return obj
        if isinstance(obj, dict):
            for key in ("results","verdicts","classifications","data"):
                if key in obj and isinstance(obj[key], list):
                    return obj[key]
            for v in obj.values():
                if isinstance(v, list):
                    return v
    except Exception:
        pass
    return _extract_json_array(text)

def ai_filter_batch(
    records: List[Dict],
    batch_size: int = 20,
    unreachable_sites=None,
    date_range=None,
    progress_callback=None,
) -> Tuple[List[Dict], Dict]:
    info: Dict = {
        "used": False, "attempted": False, "reason": "", "error_detail": "",
        "model": "", "batches": 0, "errors": 0,
        "input": len(records), "kept": 0,
        "system_prompt": "", "unreachable_addendum": "", "unreachable_count": 0,
    }
    api_key = os.getenv("GROQ_API_KEY")
    model = os.getenv("GROQ_MODEL", DEFAULT_GROQ_MODEL)

    system_prompt = AI_SYSTEM_PROMPT
    addendum_parts = []
    if date_range:
        addendum_parts.append(f"\n\nDATE RANGE OF INTEREST: {date_range[0]} → {date_range[1]}.")
    if unreachable_sites:
        info["unreachable_count"] = len(unreachable_sites)
        site_lines = "\n".join(
            f"- {s['company']} ({s.get('reason','unknown')}): {s['url']}"
            for s in unreachable_sites[:60]
        )
        addendum_parts.append(
            "\n\nUNREACHABLE PR SITES (could not be scraped this run):\n"
            + site_lines
            + "\n\nFor these companies, draw on your own knowledge if you are aware of "
            "significant announcements within the date range. Include synthetic entries "
            "with ai_synthesized: true. Never fabricate."
        )
    if addendum_parts:
        system_prompt = AI_SYSTEM_PROMPT + "".join(addendum_parts)
    info["system_prompt"] = system_prompt
    info["unreachable_addendum"] = "".join(addendum_parts)

    if not api_key:
        msg = "GROQ_API_KEY not set — AI filter skipped, using keyword filter only."
        logger.info(f"[AI] {msg}")
        info["reason"] = msg
        for r in records:
            r["ai_relevant"] = True
            r["ai_reason"] = "AI filter skipped (no API key)"
        info["kept"] = len(records)
        return records, info

    info["attempted"] = True
    info["model"] = model
    kept: List[Dict] = []
    total_batches = (len(records) + batch_size - 1) // batch_size

    for batch_idx, batch_start in enumerate(range(0, len(records), batch_size), start=1):
        batch = records[batch_start: batch_start + batch_size]
        info["batches"] += 1
        if progress_callback:
            try:
                progress_callback(batch_idx, total_batches, "running")
            except Exception:
                pass
        payload = [{"company": r.get("company",""), "headline": r.get("headline","")} for r in batch]
        user_msg = (
            f"Classify the following {len(payload)} headlines. "
            'Return a JSON object: {"results": [ ... ]} where each element is '
            '{"headline": "...", "relevant": true|false, "reason": "..."}.\n\n'
            + json.dumps(payload, ensure_ascii=False)
        )
        try:
            text = _call_groq(api_key, model, user_msg, system_prompt=system_prompt)
            verdicts = _parse_verdicts(text)
        except Exception as e:
            info["errors"] += 1
            if not info["error_detail"]:
                info["error_detail"] = str(e)[:400]
            logger.warning(f"[AI] batch {batch_idx}/{total_batches} failed; keeping records (fail-open).")
            for r in batch:
                r["ai_relevant"] = True
                r["ai_reason"] = f"AI error: {type(e).__name__}"
                kept.append(r)
            if progress_callback:
                try:
                    progress_callback(batch_idx, total_batches, "error")
                except Exception:
                    pass
            continue

        by_headline = {}
        synthesized = []
        for v in verdicts:
            if not isinstance(v, dict) or "headline" not in v:
                continue
            if v.get("ai_synthesized") is True:
                synthesized.append(v)
            else:
                by_headline[v["headline"].strip().lower()] = v

        for r in batch:
            v = by_headline.get((r.get("headline") or "").strip().lower())
            if v is None:
                r["ai_relevant"] = True
                r["ai_reason"] = "no verdict returned"
                kept.append(r)
                continue
            relevant = bool(v.get("relevant", True))
            r["ai_relevant"] = relevant
            r["ai_reason"] = str(v.get("reason",""))[:200]
            if relevant:
                kept.append(r)

        today_str = date.today().strftime("%Y-%m-%d")
        for v in synthesized:
            kept.append({
                "company": str(v.get("company","Unknown"))[:120],
                "screened_entity": str(v.get("company","Unknown"))[:120],
                "screening_category": "PR Site Screening",
                "headline": str(v.get("headline",""))[:500],
                "date": today_str, "url": "",
                "source_type": "AI Knowledge (site unreachable)",
                "ai_relevant": True, "ai_reason": str(v.get("reason","AI-synthesized"))[:200],
                "ai_synthesized": True,
            })
        if synthesized:
            info["synthesized_count"] = info.get("synthesized_count", 0) + len(synthesized)

        if progress_callback:
            try:
                progress_callback(batch_idx, total_batches, "done")
            except Exception:
                pass

    info["kept"] = len(kept)
    if info["errors"] == 0:
        info["used"] = True
        info["reason"] = f"AI filter ({model}) ran on {info['batches']} batch(es). Kept {info['kept']} of {info['input']}."
    elif info["errors"] == info["batches"]:
        info["used"] = False
        hint = ""
        d = info.get("error_detail","").lower()
        if "429" in d or "rate" in d:
            hint = " Groq rate-limit hit — wait and retry."
        elif "401" in d or "invalid api key" in d:
            hint = " Groq returned 401 — API key invalid or revoked."
        info["reason"] = f"AI filter FAILED for all {info['batches']} batches — falling back to keyword-only results." + hint
    else:
        info["used"] = True
        info["reason"] = f"AI filter partially worked: {info['errors']}/{info['batches']} batches failed."
    logger.info(f"[AI] {info['reason']}")
    return kept, info


# ============================================================
# EXPORTER
# ============================================================
OUTPUT_DIR = "output"
COLUMNS = [
    ("Serial Number",    "serial_number",       8),
    ("Company Name",     "company",             24),
    ("Screened Entity",  "screened_entity",     24),
    ("Screening Category","screening_category", 22),
    ("Date",             "date",                13),
    ("News Type",        "news_type",           18),
    ("Headline",         "headline",            60),
    ("Source Link",      "url",                 40),
    ("Source Type",      "source_type",         14),
    ("Hot vs Non-Hot",   "hot",                 14),
    ("Date Collected",   "date_collected",      14),
]
NAVY = "1B3A5C"
HOT_YELLOW = "FFF2CC"
ALT_GREY = "F5F7FA"
_THIN = Side(style="thin", color="D0D7DE")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

def _ensure_output_dir():
    os.makedirs(OUTPUT_DIR, exist_ok=True)

def _write_header(ws):
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=NAVY)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx, (label, _key, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

def _write_row(ws, row_idx, record, is_hot, alt):
    align = Alignment(vertical="center", wrap_text=True)
    fill = None
    if is_hot:
        fill = PatternFill("solid", fgColor=HOT_YELLOW)
    elif alt:
        fill = PatternFill("solid", fgColor=ALT_GREY)
    for col_idx, (_label, key, _w) in enumerate(COLUMNS, start=1):
        value = record.get(key, "")
        if key == "screened_entity" and not value:
            value = record.get("company", "")
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = align
        cell.border = _BORDER
        if fill:
            cell.fill = fill
        if key == "url" and value:
            cell.hyperlink = value
            cell.font = Font(name="Calibri", size=10, color="0563C1", underline="single")
        else:
            cell.font = Font(name="Calibri", size=10, bold=is_hot and key == "hot")
    ws.row_dimensions[row_idx].height = 45

def _write_summary(ws, records, start_date, end_date):
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_fill = PatternFill("solid", fgColor=NAVY)
    label_font = Font(name="Calibri", size=11, bold=True)
    ws["A1"] = "DNL Report — Summary"
    ws["A1"].font = title_font
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 28
    ws["A3"] = "Date Range";  ws["B3"] = f"{start_date} to {end_date}"
    ws["A4"] = "Generated";   ws["B4"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws["A5"] = "Total Records"; ws["B5"] = len(records)
    ws["A6"] = "HOT Items";   ws["B6"] = sum(1 for r in records if r.get("hot") == "HOT")
    ws["A7"] = "Product Screening Results"
    ws["B7"] = sum(1 for r in records if r.get("screening_category") == "Product Screening")
    for row in range(3, 8):
        ws.cell(row=row, column=1).font = label_font
    ws["A8"] = "News Type"; ws["B8"] = "Count"
    ws["A8"].font = label_font; ws["B8"].font = label_font
    ws["A8"].fill = PatternFill("solid", fgColor="E8EEF7")
    ws["B8"].fill = PatternFill("solid", fgColor="E8EEF7")
    counts: Dict[str, int] = {}
    for r in records:
        nt = r.get("news_type") or "Other"
        counts[nt] = counts.get(nt, 0) + 1
    row_idx = 9
    for nt, c in sorted(counts.items(), key=lambda x: -x[1]):
        ws.cell(row=row_idx, column=1, value=nt)
        ws.cell(row=row_idx, column=2, value=c)
        row_idx += 1
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16

def export_to_excel(records, start_date, end_date) -> str:
    _ensure_output_dir()
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"DNL_{start_date}_{end_date}_{ts}.xlsx"
    filepath = os.path.abspath(os.path.join(OUTPUT_DIR, filename))
    wb = Workbook()
    ws = wb.active
    ws.title = "DNL Report"
    _write_header(ws)
    for i, rec in enumerate(records):
        is_hot = rec.get("hot") == "HOT"
        _write_row(ws, row_idx=i + 2, record=rec, is_hot=is_hot, alt=(i % 2 == 1))
    last_col = get_column_letter(len(COLUMNS))
    last_row = max(1, len(records) + 1)
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    summary = wb.create_sheet("Summary")
    _write_summary(summary, records, start_date, end_date)
    wb.save(filepath)
    return filepath


# ============================================================
# PIPELINE
# ============================================================
def run_pipeline(
    start_date, end_date,
    use_ai_filter=True, use_google=True, use_pr=True,
    include_undated=True,
    google_progress_cb=None, pr_progress_cb=None, pr_failure_cb=None,
    ai_progress_cb=None, stage_cb=None,
    use_products=True, product_progress_cb=None,
) -> Tuple[List[Dict], List[Dict], Dict]:
    info: Dict = {
        "ai": {}, "stages": [],
        "screening": {
            "competitors": {"screened": 0, "raw_results": 0, "final_results": 0},
            "pr_sites":    {"screened": 0, "raw_results": 0, "final_results": 0},
            "products":    {"screened": 0, "raw_results": 0, "final_results": 0},
        },
    }

    def _stage(stage_id, state, label, detail=None):
        info["stages"].append({"id": stage_id, "state": state, "label": label, "detail": detail})
        logger.info(f"[STAGE] {stage_id} → {state} :: {label} {detail or ''}")
        if stage_cb:
            try:
                stage_cb(stage_id, state, label, detail)
            except Exception:
                pass

    logger.info(f"Pipeline started: {start_date} → {end_date}")

    raw_google: List[Dict] = []
    if use_google:
        _stage("google_fetch", "running", "Screening competitor news")
        raw_google = fetch_google_news(COMPANIES, start_date, end_date, progress_callback=google_progress_cb)
        info["screening"]["competitors"].update(screened=len(COMPANIES), raw_results=len(raw_google))
        _stage("google_fetch", "complete", f"Competitor Screening: {len(raw_google)} headlines in date range", detail=f"{len(COMPANIES)} competitors screened")
    else:
        _stage("google_fetch", "skipped", "Competitor Screening disabled")

    raw_pr: List[Dict] = []
    failures: List[Dict] = []
    if use_pr:
        _stage("pr_scrape", "running", "Scraping PR / Company websites")
        raw_pr, failures = scrape_pr_websites(
            PR_WEBSITES, start_date, end_date,
            progress_callback=pr_progress_cb, failure_callback=pr_failure_cb,
            include_undated=include_undated,
        )
        for record in raw_pr:
            record.setdefault("screened_entity", record.get("company",""))
            record.setdefault("screening_category", "PR Site Screening")
        info["screening"]["pr_sites"].update(screened=len(PR_WEBSITES), raw_results=len(raw_pr))
        _stage("pr_scrape", "complete", f"PR Websites: {len(raw_pr)} headlines", detail=f"{len(PR_WEBSITES)} sites scanned, {len(failures)} failures")
    else:
        _stage("pr_scrape", "skipped", "PR scraping disabled")

    raw_products: List[Dict] = []
    if use_products:
        _stage("product_screening", "running", "Screening product news")
        raw_products = fetch_product_news(PRODUCTS, start_date, end_date, progress_callback=product_progress_cb)
        info["screening"]["products"].update(screened=len(PRODUCTS), raw_results=len(raw_products))
        _stage("product_screening", "complete", f"Product Screening: {len(raw_products)} headlines in date range", detail=f"{len(PRODUCTS)} products screened")
    else:
        _stage("product_screening", "skipped", "Product Screening disabled")

    all_records = raw_google + raw_pr + raw_products
    _stage("date_filter", "complete", f"Date range applied: {start_date} → {end_date}", detail=f"{len(all_records)} records in range")

    _stage("keyword_filter", "running", "Applying keyword relevancy filter")
    keyword_filtered = [r for r in all_records if keyword_is_relevant(r.get("headline",""))]
    _stage("keyword_filter", "complete", f"Keyword filter: {len(keyword_filtered)} of {len(all_records)} kept")

    api_key_present = bool(os.getenv("GROQ_API_KEY"))
    if use_ai_filter and api_key_present:
        _stage("ai_filter", "running", "Running AI relevancy filter (Groq)")
        unreachable = [{"company": f["company"], "url": f["url"], "reason": f.get("category","unknown")} for f in failures]
        ai_filtered, ai_info = ai_filter_batch(
            keyword_filtered, unreachable_sites=unreachable,
            date_range=(start_date, end_date), progress_callback=ai_progress_cb,
        )
        info["ai"] = ai_info
        if ai_info.get("used"):
            _stage("ai_filter", "complete", f"AI filter: kept {ai_info['kept']} of {ai_info['input']} ({ai_info.get('model','groq')})", detail=ai_info.get("reason"))
        else:
            _stage("ai_filter", "error", "AI filter failed — fell back to keyword results", detail=ai_info.get("reason"))
    else:
        ai_filtered = keyword_filtered
        for r in ai_filtered:
            r.setdefault("ai_relevant", True)
            r.setdefault("ai_reason", "AI filter not run")
        reason = "AI filter disabled in UI" if not use_ai_filter else "GROQ_API_KEY not set — keyword-only filtering"
        info["ai"] = {
            "used": False, "attempted": False, "reason": reason,
            "model": "", "batches": 0, "errors": 0,
            "input": len(keyword_filtered), "kept": len(keyword_filtered),
        }
        _stage("ai_filter", "skipped", reason)

    _stage("dedupe", "running", "Removing duplicate headlines")
    deduped = deduplicate(ai_filtered)
    _stage("dedupe", "complete", f"Dedupe: {len(deduped)} unique records ({len(ai_filtered) - len(deduped)} duplicates removed)")

    _stage("classify", "running", "Classifying news types & HOT tags")
    today_str = datetime.now().strftime("%Y-%m-%d")
    hot_count = 0
    for idx, r in enumerate(deduped, start=1):
        head = r.get("headline","")
        r["news_type"] = classify_news_type(head)
        r["hot"] = classify_hot(head)
        if r["hot"] == "HOT":
            hot_count += 1
        r["serial_number"] = idx
        r["date_collected"] = today_str
    _stage("classify", "complete", f"Classified {len(deduped)} records ({hot_count} HOT)")

    final_counts = {"Competitor Screening": 0, "PR Site Screening": 0, "Product Screening": 0}
    for record in deduped:
        cat = record.get("screening_category")
        if cat in final_counts:
            final_counts[cat] += 1
    info["screening"]["competitors"]["final_results"] = final_counts["Competitor Screening"]
    info["screening"]["pr_sites"]["final_results"]    = final_counts["PR Site Screening"]
    info["screening"]["products"]["final_results"]    = final_counts["Product Screening"]

    logger.info(f"Pipeline complete: {len(deduped)} final records")
    return deduped, failures, info


# ============================================================
# STREAMLIT UI
# ============================================================
_AUTH_EMAIL    = os.getenv("DNL_LOGIN_EMAIL",    "dnladmin@automation.in")
_AUTH_PASSWORD = os.getenv("DNL_LOGIN_PASSWORD", "123456789")

_page = st.query_params.get("page", "")

# ---- Route: dashboard shell (default) ----
if _page != "research-setup":
    st.set_page_config(
        page_title="Competitive Intelligence Engine",
        page_icon="🧭",
        layout="wide",
        initial_sidebar_state="collapsed",
    )
    st.markdown("""
    <style>
      #MainMenu, header, footer { visibility: hidden; height: 0; }
      [data-testid="stSidebar"], [data-testid="stSidebarNav"],
      [data-testid="collapsedControl"] { display: none !important; }
      .block-container { padding: 0 !important; max-width: 100% !important; }
      html, body, [data-testid="stAppViewContainer"], [data-testid="stMain"] {
        background: #0F0A2E !important; margin: 0 !important; padding: 0 !important;
      }
      iframe { border: none !important; }
    </style>
    """, unsafe_allow_html=True)

    _HTML_PATH = Path(__file__).parent / "index.html"
    try:
        _html = _HTML_PATH.read_text(encoding="utf-8")
    except FileNotFoundError:
        st.error(f"Dashboard file not found: {_HTML_PATH}")
        st.stop()

    _html = _html.replace("__PRODUCT_COUNT__", str(len(PRODUCTS)))
    _html = _html.replace("__COMPETITOR_COUNT__", str(len(COMPANIES)))
    components.html(_html, height=1400, scrolling=True)
    st.stop()


# ---- Route: DNL pipeline UI ----
st.set_page_config(page_title="DNL Automation Tool", page_icon="📰", layout="wide")

if "auth_ok" not in st.session_state:
    st.session_state.auth_ok = False

def _render_login():
    st.markdown("""
    <div style='max-width:420px;margin:80px auto 24px auto;padding:32px;
    border-radius:12px;background:linear-gradient(135deg,#1B3A5C 0%,#2A5A8C 100%);
    color:#fff;text-align:center'>
    <h2 style='margin:0 0 8px 0;color:#fff'>📰 DNL Automation</h2>
    <div style='font-size:13px;opacity:0.9'>Sign in to continue</div>
    </div>
    """, unsafe_allow_html=True)
    col_l, col_c, col_r = st.columns([1, 2, 1])
    with col_c:
        with st.form("login_form", clear_on_submit=False):
            email    = st.text_input("Email", placeholder="you@company.com")
            password = st.text_input("Password", type="password")
            submitted = st.form_submit_button("🔓 Sign in", use_container_width=True, type="primary")
        if submitted:
            if email.strip().lower() == _AUTH_EMAIL.lower() and password == _AUTH_PASSWORD:
                st.session_state.auth_ok = True
                st.rerun()
            else:
                st.error("Invalid email or password.")

if not st.session_state.auth_ok:
    _render_login()
    st.stop()

# ---- Theme ----
st.markdown("""
<style>
  .dnl-hero { background:linear-gradient(135deg,#1B3A5C 0%,#2A5A8C 100%);
    padding:22px 28px; border-radius:10px; color:#fff; margin-bottom:18px; }
  .dnl-hero h1 { margin:0; font-size:28px; }
  .hot-metric div[data-testid="stMetricValue"] { color:#C62828; }
  .stage-row { padding:10px 14px; border-radius:8px; margin:4px 0;
    border:1px solid #E1E6ED; background:#FAFBFC;
    display:flex; justify-content:space-between; align-items:center; font-size:14px; color:#1B1B1B !important; }
  .stage-row * { color:#1B1B1B !important; }
  .stage-running  { background:#FFF3C4 !important; border-color:#E0A800 !important; }
  .stage-complete { background:#CDEBD3 !important; border-color:#2E7D32 !important; }
  .stage-skipped  { background:#E6E6E6 !important; border-color:#9E9E9E !important; }
  .stage-error    { background:#F8C9C9 !important; border-color:#B71C1C !important; }
  .stage-label    { font-weight:700; color:#1B1B1B !important; }
  .stage-detail   { font-size:12px; color:#3D3D3D !important; margin-top:2px; }
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="dnl-hero"><h1>DNL Automation Tool</h1></div>', unsafe_allow_html=True)

# Back to Dashboard link
st.markdown("""
<a href="./" target="_self"
   style="position:fixed;top:12px;right:16px;z-index:9999;
          background:#6D4FF6;color:#fff;padding:8px 14px;
          border-radius:8px;font-size:13px;font-weight:600;text-decoration:none;">
  ← Back to Dashboard
</a>
""", unsafe_allow_html=True)

# ---- Sidebar ----
with st.sidebar:
    st.caption(f"👤 Signed in as `{_AUTH_EMAIL}`")
    if st.button("🚪 Sign out", use_container_width=True):
        st.session_state.auth_ok = False
        st.rerun()
    st.divider()
    st.subheader("Date Range")
    default_start = date.today() - timedelta(days=3)
    default_end   = date.today()
    start_date = st.date_input("Start date", value=default_start)
    end_date   = st.date_input("End date",   value=default_end)
    st.subheader("Sources")
    use_google   = st.checkbox("Competitor Screening (Google News RSS)", value=True)
    use_pr       = st.checkbox("PR / Company Websites", value=True)
    use_products = st.checkbox("Product Screening", value=bool(PRODUCTS), disabled=not PRODUCTS)
    if not PRODUCTS:
        st.caption("Add products to PRODUCT_CATALOG in main.py to enable this source.")
    include_undated = st.checkbox("Include PR articles without a detectable date", value=True)
    st.subheader("AI Filtering")
    use_ai  = st.checkbox("Enable AI Relevancy Filter (Groq)", value=True)
    has_key = bool(os.getenv("GROQ_API_KEY"))
    if use_ai and not has_key:
        st.warning("⚠️ `GROQ_API_KEY` not detected in `.env`. AI filter will be skipped.")
    elif use_ai and has_key:
        st.caption(f"✓ Groq key detected. Model: `{os.getenv('GROQ_MODEL', 'llama-3.3-70b-versatile')}`")

# ---- Session state ----
st.session_state.setdefault("records", None)
st.session_state.setdefault("failures", [])
st.session_state.setdefault("excel_path", None)
st.session_state.setdefault("pipeline_info", None)

# ---- Stage board ----
STAGES = [
    ("google_fetch",       "🏢  Competitor Screening"),
    ("pr_scrape",          "🌐  PR website scrape"),
    ("product_screening",  "📦  Product Screening"),
    ("date_filter",        "📅  Date range filter"),
    ("keyword_filter",     "🔑  Keyword filter"),
    ("ai_filter",          "🤖  AI relevancy filter (Groq)"),
    ("dedupe",             "🧹  Deduplicate"),
    ("classify",           "🏷️  Classify (HOT + News Type)"),
]
STATE_ICON = {"pending":"⏳","running":"⏱️","complete":"✅","skipped":"⏭️","error":"❌"}

def render_stage_board(stage_state, placeholder):
    rows = []
    for stage_id, default_label in STAGES:
        info = stage_state.get(stage_id, {"state":"pending","label":default_label,"detail":None})
        state = info.get("state","pending")
        label = info.get("label") or default_label
        detail = info.get("detail") or ""
        icon = STATE_ICON.get(state,"•")
        css_class = f"stage-{state}" if state in ("running","complete","skipped","error") else ""
        rows.append(
            f'<div class="stage-row {css_class}">'
            f'<div><span style="font-size:18px;margin-right:8px">{icon}</span>'
            f'<span class="stage-label">{label}</span>'
            f'{("<div class=stage-detail>" + detail + "</div>") if detail else ""}'
            f'</div></div>'
        )
    placeholder.markdown("\n".join(rows), unsafe_allow_html=True)

# ---- Tabs ----
run_clicked = st.button("🚀 Run Newsletter Scan", use_container_width=True, type="primary")
tab_scan, tab_hub = st.tabs(["📊 Scan & Results", "📚 Knowledge Hub"])

with tab_scan:
    if run_clicked:
        if start_date > end_date:
            st.error("Start date must be on or before end date.")
        elif not (use_google or use_pr or use_products):
            st.error("Select at least one screening source.")
        else:
            gcol, pcol, product_col = st.columns(3)
            with gcol:
                st.markdown("**🏢 Competitor Screening progress**")
                g_bar = st.progress(0.0, text="Idle" if not use_google else "Waiting...")
            with pcol:
                st.markdown("**🟩 PR Websites progress**")
                p_bar = st.progress(0.0, text="Idle" if not use_pr else "Waiting...")
                pr_fail_placeholder = st.empty()
                pr_failures_live: list = []

                def _render_pr_failures():
                    if not pr_failures_live:
                        pr_fail_placeholder.empty()
                        return
                    df_fail = pd.DataFrame([{
                        "Company": f.get("company",""), "Issue": f.get("category","error"),
                        "URL": f.get("url",""), "Detail": (f.get("error","") or "")[:200],
                    } for f in pr_failures_live])
                    with pr_fail_placeholder.container():
                        st.markdown(f"**⚠️ Live failures — {len(pr_failures_live)} sites failed**")
                        st.dataframe(df_fail, use_container_width=True, hide_index=True,
                            height=min(300, 45 + 35 * len(df_fail)),
                            column_config={"URL": st.column_config.LinkColumn("URL", width="medium")},
                        )

            with product_col:
                st.markdown("**📦 Product Screening progress**")
                product_bar = st.progress(0.0, text="Idle" if not use_products else "Waiting...")

            st.markdown("### Pipeline stages")
            stage_placeholder = st.empty()
            stage_state: dict = {}
            render_stage_board(stage_state, stage_placeholder)
            st.markdown("**🤖 AI filter progress (Groq batches)**")
            ai_bar = st.progress(0.0, text="Idle — waiting for previous stages...")

            def google_cb(company, idx, total):
                g_bar.progress(min(idx / max(total,1), 1.0), text=f"{idx}/{total} — {company}")

            def pr_cb(company, idx, total):
                p_bar.progress(min(idx / max(total,1), 1.0), text=f"{idx}/{total} — {company}")

            def product_cb(product, idx, total):
                product_bar.progress(min(idx / max(total,1), 1.0), text=f"{idx}/{total} — {product}")

            def pr_fail_cb(failure: dict):
                pr_failures_live.append(failure)
                _render_pr_failures()

            def ai_cb(batch_idx, total, status):
                if status == "running":
                    frac = (batch_idx - 1) / max(total,1)
                    ai_bar.progress(min(frac, 1.0), text=f"🤖 Batch {batch_idx}/{total} — calling Groq...")
                else:
                    frac = batch_idx / max(total,1)
                    icon = "✅" if status == "done" else "⚠️"
                    ai_bar.progress(min(frac, 1.0), text=f"{icon} Batch {batch_idx}/{total} {status}")

            def stage_cb(stage_id, state, label, detail):
                stage_state[stage_id] = {"state": state, "label": label, "detail": detail}
                render_stage_board(stage_state, stage_placeholder)

            try:
                records, failures, info = run_pipeline(
                    start_date=start_date, end_date=end_date,
                    use_ai_filter=use_ai, use_google=use_google,
                    use_pr=use_pr, use_products=use_products,
                    include_undated=include_undated,
                    google_progress_cb=google_cb, pr_progress_cb=pr_cb,
                    product_progress_cb=product_cb, pr_failure_cb=pr_fail_cb,
                    ai_progress_cb=ai_cb, stage_cb=stage_cb,
                )
                if use_google:   g_bar.progress(1.0, text="✅ Competitor Screening done")
                if use_pr:       p_bar.progress(1.0, text="✅ PR Websites done")
                if use_products: product_bar.progress(1.0, text="✅ Product Screening done")
                ai_info_final = info.get("ai", {})
                if ai_info_final.get("attempted"):
                    ai_bar.progress(1.0, text=f"✅ AI done — {ai_info_final.get('batches',0)} batches, {ai_info_final.get('errors',0)} errors, kept {ai_info_final.get('kept',0)}/{ai_info_final.get('input',0)}")
                else:
                    ai_bar.progress(1.0, text="⏭️ AI skipped")

                st.session_state.records = records
                st.session_state.failures = failures
                st.session_state.pipeline_info = info

                ai_info = info.get("ai", {})
                if ai_info.get("used"):
                    st.success(f"🤖 {ai_info.get('reason','AI filter ran')}")
                elif ai_info.get("attempted"):
                    st.error(f"⚠️ AI didn't work — switching to keywords. {ai_info.get('reason','')}")
                else:
                    st.info(f"ℹ️ {ai_info.get('reason','AI filter not run')}")

                if ai_info.get("system_prompt"):
                    with st.expander(f"🔍 View AI prompt sent to Groq (model: {ai_info.get('model','—')}, unreachable: {ai_info.get('unreachable_count',0)}, synthesized: {ai_info.get('synthesized_count',0)})", expanded=False):
                        st.code(ai_info["system_prompt"], language="markdown")

                if records:
                    excel_path = export_to_excel(records, start_date, end_date)
                    st.session_state.excel_path = excel_path
                    st.success(f"Scan complete: {len(records)} records exported.")
                else:
                    st.session_state.excel_path = None
                    st.warning("No relevant records found for the selected range.")
            except Exception as e:
                st.error(f"Pipeline failed: {e}")
                st.exception(e)

    records = st.session_state.records
    failures = st.session_state.failures
    excel_path = st.session_state.excel_path

    if records is not None:
        total = len(records)
        hot = sum(1 for r in records if r.get("hot") == "HOT")
        p_count = sum(1 for r in records if r.get("source_type") == "PR Website")
        pipeline_info = st.session_state.pipeline_info or {}
        screening_info = pipeline_info.get("screening", {})
        competitor_stats = screening_info.get("competitors", {})
        product_stats = screening_info.get("products", {})
        g_count = sum(1 for r in records if r.get("source_type") == "Google News")

        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Total Records", total)
        with c2:
            st.markdown('<div class="hot-metric">', unsafe_allow_html=True)
            st.metric("HOT Items", hot)
            st.markdown("</div>", unsafe_allow_html=True)
        c3.metric("Google News Results", g_count)
        c4.metric("PR Websites", p_count)

        s1, s2, s3, s4 = st.columns(4)
        s1.metric("Competitors Screened", competitor_stats.get("screened",0))
        s2.metric("Competitor Results", competitor_stats.get("final_results",0))
        s3.metric("Products Screened", product_stats.get("screened",0))
        s4.metric("Product Results", product_stats.get("final_results",0))

        if excel_path and os.path.exists(excel_path):
            with open(excel_path, "rb") as f:
                st.download_button("⬇️ Download Excel Report", data=f.read(),
                    file_name=os.path.basename(excel_path),
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )

        if total > 0:
            st.subheader("Preview")
            fc1, fc2, fc3, fc4 = st.columns(4)
            hot_filter    = fc1.selectbox("Hot filter", ["All","HOT only","Non-Hot only"])
            type_opts     = ["All"] + sorted({r.get("news_type","Other") for r in records})
            type_filter   = fc2.selectbox("News Type", type_opts)
            source_opts   = ["All"] + sorted({r.get("source_type","") for r in records})
            source_filter = fc3.selectbox("Source Type", source_opts)
            cat_opts      = ["All"] + sorted({r.get("screening_category","Uncategorized") for r in records})
            cat_filter    = fc4.selectbox("Screening Category", cat_opts)

            filtered = records
            if hot_filter == "HOT only":   filtered = [r for r in filtered if r.get("hot") == "HOT"]
            elif hot_filter == "Non-Hot only": filtered = [r for r in filtered if r.get("hot") != "HOT"]
            if type_filter   != "All": filtered = [r for r in filtered if r.get("news_type") == type_filter]
            if source_filter != "All": filtered = [r for r in filtered if r.get("source_type") == source_filter]
            if cat_filter    != "All": filtered = [r for r in filtered if r.get("screening_category","Uncategorized") == cat_filter]

            df = pd.DataFrame([{
                "#": r.get("serial_number"),
                "Screening Category": r.get("screening_category",""),
                "Company": r.get("company"),
                "Screened Entity": r.get("screened_entity") or r.get("company"),
                "Date": r.get("date"),
                "Type": r.get("news_type"),
                "Headline": (r.get("headline") or "")[:120],
                "Source": r.get("source_type"),
                "Hot": r.get("hot"),
            } for r in filtered])
            st.dataframe(df, height=500, use_container_width=True, hide_index=True)
            st.caption(f"Showing {len(filtered)} of {total} records.")

        if failures:
            with st.expander(f"⚠️ {len(failures)} scraping failures"):
                for f in failures:
                    st.markdown(f"**{f.get('company')}** — `{f.get('category','error')}`  \n{f.get('url')}  \n_{f.get('error')}_")


with tab_hub:
    st.markdown("Use this view to **audit the relevancy criteria**. Every list below drives what the scanner picks up. If something important is missing or wrong, edit `main.py` and restart.")
    h1, h2, h3, h4, h5 = st.columns(5)
    h1.metric("Companies tracked", len(COMPANIES))
    h2.metric("PR websites", len(PR_WEBSITES))
    h3.metric("Products tracked", len(PRODUCTS))
    h4.metric("Relevant keywords", len(RELEVANT_KEYWORDS))
    h5.metric("HOT triggers", len(HOT_KEYWORDS))
    st.divider()

    with st.expander(f"🏢 Monitored Companies ({len(COMPANIES)})", expanded=False):
        st.dataframe(pd.DataFrame({"#": list(range(1, len(COMPANIES)+1)), "Company": COMPANIES}), hide_index=True, use_container_width=True, height=420)

    with st.expander(f"📦 Monitored Products ({len(PRODUCTS)})", expanded=False):
        if PRODUCTS:
            st.dataframe(pd.DataFrame([{"#": idx, "Company": p["company"], "Product": p["name"]} for idx, p in enumerate(PRODUCTS, start=1)]), hide_index=True, use_container_width=True, height=420)
        else:
            st.info("No products configured. Add entries to PRODUCT_CATALOG in main.py.")

    with st.expander(f"🌐 PR / Company Websites ({len(PR_WEBSITES)})", expanded=False):
        pr_df = pd.DataFrame([{"Company": s["company"], "URL": s["url"]} for s in PR_WEBSITES])
        st.dataframe(pr_df, hide_index=True, use_container_width=True, height=420,
            column_config={"URL": st.column_config.LinkColumn("URL", display_text="Open ↗")})

    with st.expander(f"✅ Relevant Keywords ({len(RELEVANT_KEYWORDS)})", expanded=False):
        kw_source = "keywords.txt" if (_kw_file and _kw_file[0]) else "main.py defaults"
        st.caption(f"Source: {kw_source}. A headline must contain at least one of these terms to pass the keyword filter.")
        st.markdown(" ".join(f"<span style='background:#E8EEF7;border:1px solid #C5D2E5;border-radius:14px;padding:3px 10px;margin:3px;display:inline-block;font-size:12px;color:#1B3A5C'>{kw}</span>" for kw in RELEVANT_KEYWORDS), unsafe_allow_html=True)

    with st.expander(f"❌ Exclude Keywords ({len(EXCLUDE_KEYWORDS)})", expanded=False):
        st.markdown(" ".join(f"<span style='background:#FDECEC;border:1px solid #F5B7B1;border-radius:14px;padding:3px 10px;margin:3px;display:inline-block;font-size:12px;color:#922B21'>{kw}</span>" for kw in EXCLUDE_KEYWORDS), unsafe_allow_html=True)

    with st.expander(f"🔥 HOT Triggers ({len(HOT_KEYWORDS)})", expanded=False):
        st.markdown(" ".join(f"<span style='background:#FFF2CC;border:1px solid #E8C547;border-radius:14px;padding:3px 10px;margin:3px;display:inline-block;font-size:12px;color:#7E5A00'>{kw}</span>" for kw in HOT_KEYWORDS), unsafe_allow_html=True)

    with st.expander(f"🏷️ News Type Rules ({len(NEWS_TYPE_RULES)} categories)", expanded=False):
        for category, triggers in NEWS_TYPE_RULES.items():
            st.markdown(f"**{category}**")
            st.markdown(" ".join(f"<span style='background:#F0F4F8;border:1px solid #C5D2E5;border-radius:12px;padding:2px 8px;margin:2px;display:inline-block;font-size:11px;color:#1B3A5C'>{t}</span>" for t in triggers), unsafe_allow_html=True)
