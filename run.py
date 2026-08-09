"""
DNL Automation Tool — headless GitHub Actions runner.

Usage (locally or in CI):
    python run.py                          # last 3 days
    python run.py --start 2025-08-01 --end 2025-08-07
    python run.py --days 7
    python run.py --no-ai                  # skip Groq, keyword-only
    python run.py --no-google              # skip Google News
    python run.py --no-pr                  # skip PR site scraping
    python run.py --no-products            # skip product screening

Output:
    output/DNL_<start>_<end>_<timestamp>.xlsx   — main report
    output/DNL_<start>_<end>_<timestamp>_failures.csv — scraping failures
    output/run_summary.md                        — human-readable summary
                                                   (used by the GitHub Actions job summary)

Environment variables (set as GitHub Secrets):
    GROQ_API_KEY          — Groq API key for AI relevancy filter
    GROQ_MODEL            — optional model override (default: llama-3.3-70b-versatile)
"""

# ── stdlib ────────────────────────────────────────────────────────────────────
import argparse
import json
import logging
import os
import re
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple, Union
from urllib.parse import quote_plus, urljoin, urlparse

# ── third-party ───────────────────────────────────────────────────────────────
import feedparser
import pandas as pd
import requests
from bs4 import BeautifulSoup
from dateutil import parser as dateparser
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

load_dotenv()

# ══════════════════════════════════════════════════════════════════════════════
# LOGGER
# ══════════════════════════════════════════════════════════════════════════════
LOG_DIR = "logs"
os.makedirs(LOG_DIR, exist_ok=True)
_log_path = os.path.join(LOG_DIR, f"dnl_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log")

logger = logging.getLogger("dnl")
logger.setLevel(logging.INFO)
if not logger.handlers:
    fmt = logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s", datefmt="%H:%M:%S")
    fh = logging.FileHandler(_log_path, encoding="utf-8")
    fh.setFormatter(fmt)
    logger.addHandler(fh)
    ch = logging.StreamHandler(sys.stdout)
    ch.setFormatter(fmt)
    logger.addHandler(ch)
    logger.propagate = False

def _log(msg: str):
    logger.info(msg)
    # Also echo to GitHub Actions step summary via GITHUB_STEP_SUMMARY if available
    gss = os.environ.get("GITHUB_STEP_SUMMARY")
    if gss:
        try:
            with open(gss, "a", encoding="utf-8") as f:
                f.write(msg + "\n")
        except Exception:
            pass


# ══════════════════════════════════════════════════════════════════════════════
# KEYWORDS  (loaded from keywords.txt if present, else built-in defaults)
# ══════════════════════════════════════════════════════════════════════════════
def _load_keywords():
    kw_path = Path("keywords.txt")
    relevant, exclude, hot = [], [], []
    if kw_path.exists():
        section = None
        for line in kw_path.read_text(encoding="utf-8").splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            upper = stripped.upper()
            if stripped.startswith("#"):
                if "RELEVANT" in upper:
                    section = "relevant"
                elif "EXCLUDE" in upper:
                    section = "exclude"
                elif "HOT" in upper:
                    section = "hot"
                continue
            if section == "relevant":
                relevant.append(stripped)
            elif section == "exclude":
                exclude.append(stripped)
            elif section == "hot":
                hot.append(stripped)
        _log(f"[Keywords] Loaded from keywords.txt — relevant:{len(relevant)} exclude:{len(exclude)} hot:{len(hot)}")

    RELEVANT_KEYWORDS = relevant or [
        "immunodiagnostic","allergy diagnostic","autoimmune diagnostic",
        "immunoassay","diagnostic platform","point-of-care","poc","ivd",
        "elisa","lateral flow","multiplex","serology",
        "allergy immunotherapy","allergen immunotherapy","scit","slit",
        "allergy therapeutic","allergy","autoimmune","immunology",
        "merger","acquisition","acquires","acquired","takeover","divest",
        "partnership","collaboration","agreement","alliance","joint venture",
        "product launch","new product","launches","launch","platform",
        "ce mark","ce-mark","fda clearance","fda approval","510(k)",
        "regulatory approval","regulatory","ema approval","market authorization",
        "conference","symposium","congress","eaaci","aaaai","acaai","escmid",
        "diagnostics","diagnostic","assay","reagent","analyzer",
        "laboratory","lab","clinical","test kit",
        "revenue","earnings","financial results",
        "ceo","cfo","appoints","appointment","leadership","executive",
    ]
    EXCLUDE_KEYWORDS = exclude or [
        "generic drug","biosimilar","chemotherapy","oncology drug",
        "cardiovascular drug","diabetes drug","obesity drug",
        "mental health","psychiatry","neurology drug","orthopedic",
        "dermatology drug","ophthalmology drug",
    ]
    HOT_KEYWORDS = hot or [
        "merger","acquisition","acquires","acquired","takeover","divest",
        "fda approval","fda clearance","ce mark","510(k) clearance",
        "market authorization","regulatory approval",
        "strategic alliance","joint venture","major partnership",
        "platform launch","breakthrough","pivotal","landmark",
    ]
    return RELEVANT_KEYWORDS, EXCLUDE_KEYWORDS, HOT_KEYWORDS

RELEVANT_KEYWORDS, EXCLUDE_KEYWORDS, HOT_KEYWORDS = _load_keywords()

NEWS_TYPE_RULES = {
    "M&A":           ["merger","acquisition","acquires","acquired","takeover","divest"],
    "Regulatory":    ["regulatory","approval","cleared","clearance","authorized","authorization","fda","ema","ce mark","ce-mark","510(k)"],
    "Product Launch":["launch","launches","new product","new platform","introduces","unveiled","unveils"],
    "Partnership":   ["partnership","collaboration","agreement","alliance","joint venture","co-develop","license","licensing"],
    "Conference":    ["conference","congress","symposium","expo","eaaci","aaaai","acaai","ecp","esh","escmid"],
    "Organizational":["ceo","cfo","appoints","appointment","leadership","executive","board","director","restructure"],
    "Financial":     ["revenue","earnings","results","financial","profit","growth","guidance"],
}

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
}
REQUEST_TIMEOUT = 15


# ══════════════════════════════════════════════════════════════════════════════
# COMPANIES, PRODUCTS, PR SITES
# ══════════════════════════════════════════════════════════════════════════════
COMPANIES = [
    "A Menarini Diagnostics","Abbott","AC Immune","Aesku Diagnostics",
    "aimmune therapeutics","Alcor Scientific","Alerje","Align ENT + Allergy",
    "AliveDX","ALK","Alladapt Immunotherapeutics","Allegra","AllerGenis",
    "Allergy Amulet","Allergy Partners","Allergy therapeutics","Allermi",
    "AllerSmart","Allervie Health","Aravax","Asit biotech","ATANIS Biotech AG",
    "Augurex Life Sciences","Autobio Diagnostics","Autu Bio",
    "Axim Biotechnologies","Bayer Healthcare LLC","BBI Solutions",
    "BD Biosciences","Beckman Coulter","Bio-Rad","Biosynex",
    "Bio-Techne Corporation","Bredis Healthcare","BSBE Biotech","Buhlmann",
    "BÜHLMANN Laboratories","C2N Diagnostics","Camallergy",
    "Canon Medical Systems","CellaVision","Celltrion","Cerascreen",
    "CF PharmTech","Circular Genomics","Cizzle Biotechnology",
    "C-Luminary Biotech","CND Life Sciences","Creative Proteomics",
    "Crown Bioscience","CVS Health","Cytox","Danaher Corporation",
    "DBV technologies","Desentum","Diadem","DiagnaMed","DiamiR","DiaSorin",
    "Differentia Biotech","DOTS Technology Corp","Dr. Fooke","Droege Group",
    "Durin Technologies","Eiken Chemical","EKF Diagnostics","empowerDx",
    "Enterome","Euro Diagnostica","Euroimmun","Eurospital","Evotec SE",
    "Exagen","Excellergy","FAES Farma S.A","Fujirebio",
    "Gentian Diagnostics","Global Medical Technologies",
    "Gold Standard Diagnostics","Grifols","HOB Biotech","Hycor",
    "HYCOR Biomedical","Iason","IgGenix","Imcyse","Immuno Concepts",
    "Immuno Diagnostic System","InBio","Infinity BiologiX","Inimmune",
    "Inmune Bio","Inova Diagnostics","Intrommune Therapeutics","Invitros",
    "JSR Life Sciences","Jubilant Pharmova","Kenota Health","LG Chem",
    "MADX","MannKind","MBL","MedicalSystem Biotechnology",
    "Mediwiss Analytics GmbH","Medtronic","Menarini","Merck",
    "Meso Scale Diagnostics","Mikrogen","Mindray","Morepen",
    "MYND Diagnostics","Nanjing Vazyme Biotech","Neogen","Nippon Chemiphar",
    "Octave Bioscience","Omega Diagnostics","Oncimmune","Orgentec",
    "Ortho Clinical Diagnostics","PEPperPRINT GmbH","Perspectum",
    "PrecisionLife","Progentec Diagnostics, Inc.","Prolight Diagnostics",
    "Prometheus Laboratories","Quidel","QuidelOrtho","Quotient","R-Biopharm",
    "Reacta Biotech","Resonac","Revvity","Roche Diagnostics","Sebia",
    "Siemens Healthineers","Snibe","Stallergenes Greer","Sugentech",
    "Svar Lifesciences","Theradiag","Trinity Biotech","Ukko","Werfen",
    "WuXi Diagnostics","YHLO","ZEUS Scientific",
]

PRODUCT_CATALOG = {
    "A. Menarini Diagnostics": ["ZENIT","Zenit PRO","Zenit fast","Zenit Reagents","Zenit HUB","Zenit Prime","CHORUS","CHORUS TRIO"],
    "Abbott":                  ["Alinity i","ARCHITECT","ARCHITECT i1000SR","ARCHITECT i2000SR","Alinity ci-series","CHEMIFLEX","Lab Central"],
    "AESKU.GROUP":             ["AESKULISA","AESKUSLIDES","AESKUBLOTS","AESQC","HELIOS","AKLIDES","akiron","DST","ATLAS"],
    "AliveDx":                 ["MosaiQ","LumiQ","Alba"],
    "Autobio":                 ["AutoLumo","AutoLumo A1000","AutoLumo A1860","AutoLumo A2000","AutoLumo A2000 Plus","AutoLumo A6200","AutoLumo S900","Autobio CLIA","CLIA Microparticles"],
    "Beckman Coulter":         ["Access","Access 2","DxI","DxI 9000","DxI 9000 Access","UniCel DxI","UniCel DxC","DxC 500i"],
    "Bio-Rad":                 ["BioPlex 2200","BioPlex 2200 System","BioPlex 2200 Autoimmune Panels","Kallestad","EVOLIS","Platelia","BioPlex"],
    "BUHLMANN Laboratories":   ["CAST","CAST ELISA","Flow CAST","GanglioCombi","Quantum Blue","BÜHLMANN fCAL","fCAL turbo","IBDoc","Calex","CalApp"],
    "C-Luminary Biotech":      ["Sharay","Sharay 1600","Sharay 6000","Sharay 8000","Sharay 800","C-Luminary CLIA"],
    "DiaSorin":                ["LIAISON","LIAISON XL","LIAISON XS","LIAISON IQ","xMAP","xMAP INTELLIFLEX","Simplexa","Luminex","VERIGENE"],
    "Dr. Fooke":               ["ALLERG-O-LIQ","ALLERG-O-SYSTEM","ALFA","ALFA Reader","ALLERG-O-WIN","ARTHUS","LFA Reader"],
    "Eurospital":              ["Eu-tTG","alpha-GliaPep","alpha-Gliatest","Xeliac","Xeliac Test Pro","XeliGen","EAL","ECD","Easy-NAT","Calprest"],
    "Gold Standard Diagnostics":["Hurricane Dx","AESKULISA","NovaLisa","The Bolt","ThunderBolt","AIX1000","NovaTec","Virotech"],
    "HOB Biotech":             ["BioCLIA","BioCLIA 500","BioCLIA 6500","BioCLIA Allergy","BioCLIA Autoimmune","BioLINE","BioCLIA 1200","REAST"],
    "HYCOR":                   ["NOVEOS","HYTEC","HYTEC 288 Plus","NOVEOS Allergens"],
    "Invitros":                ["AdvanSure AlloScreen","AdvanSure","AdvanSure i3 TB-IGRA","AdvanSure RV-Plus","AdvanSure E3 System"],
    "MADx":                    ["ALEX","ALEX2","ALEX3","Allergy Xplorer","FOX","Food Xplorer","ImageXplorer","MAX 9k","MAX45k","Raptor Server","MADx Cloud"],
    "MBL Life Science":        ["MESACUP","MESACUP Series","MEBLux","STACIA","MBL IVD","MBL Autoimmune"],
    "Nippon Chemiphar":        ["DropScreen","DropScreen ST-1","Drop Screen","IgE NC","R-NanoBio"],
    "QuidelOrtho":             ["MICROVUE","MICROVUE Complement Multiplex","THYRETAIN","VITROS","ELISA Kits","TECOMedical","SPG"],
    "Revvity / EUROIMMUN / IDS":["EUROLINE","EUROBlotOne","EUROLineScan","EUROPattern","IDS-iSYS","IDS-i10","IDS-i20","EUROLabWorkstation","EUROLabOffice 4.0","EUROIMMUN Analyzer I","IF Sprinter","Sprinter XL"],
    "Roche":                   ["Elecsys","cobas e","cobas e 411","cobas e 402","cobas e 801","cobas pure","cobas pro","ECLIA"],
    "Sebia":                   ["Alegria","Alegria 2","Alegria Monotest","SMC Technology","ORGENTEC","myAlegria"],
    "Siemens Healthineers":    ["IMMULITE","IMMULITE 2000","IMMULITE 2000 XPi","IMMULITE 1000","3gAllergy","IMMULITE 3gAllergy","Atellica","Atellica Solution","Atellica CI","ADVIA Centaur","Dimension Vista","Dimension EXL"],
    "Snibe Diagnostics":       ["MAGLUMI","MAGLUMI X3","MAGLUMI X6","MAGLUMI X8","MAGLUMI 600","MAGLUMI 800","MAGLUMI 2000","MAGLUMI 2000 Plus","MAGLUMI M Series","Biolumi CX8","X-TECH"],
    "Sugentech":               ["SGTi-Allergy Screen","SGTi-Allergy Screen PLUS","SGTi-Allergy PLUS 60","S-Blot 3","S-Blot 3 PLUS","S-Blot 2","S-Blot 2 Easy","S-Blot 2 Easy PLUS","INCLIX F-100","INCLIX"],
    "Theradiag":               ["FIDIS","Lisa Tracker","LISA TRACKER","i-Track10","i-Tracker","xMAP","Theranostics","Biosynex","Luminex"],
    "Trinity Biotech":         ["ImmuGlo","Mardx","ImmuLisa","Captia","ImmuBlot","Marstripe","ImmcoStripe","Amerlex","Immco","Immco Diagnostics","Trinity Autoimmune"],
    "Werfen":                  ["Aptiva","BIO-FLASH","QUANTA Flash","QUANTA Lite","NOVA View","NOVA Lite","QUANTA-Lyser 3000","QUANTA Link","Integrated Lab+","Inova Diagnostics"],
    "YHLO":                    ["iFlash","iFlash 1200","iFlash 1800","iFlash 3000","iFlash 9000","iStar 500","CLIA Solution","LIA Solution"],
}
PRODUCTS = [{"company": c, "name": p} for c, prods in PRODUCT_CATALOG.items() for p in prods]

PR_WEBSITES = [
    {"company": "Aesku Diagnostics",      "url": "https://www.aesku.com/"},
    {"company": "Abbott",                  "url": "https://abbott.mediaroom.com/press-releases"},
    {"company": "Danaher",                 "url": "https://www.danaher.com/newsroom"},
    {"company": "Beckman Coulter",         "url": "https://www.beckmancoulter.com/about-beckman-coulter/newsroom/press-releases"},
    {"company": "Bio-Rad (IR)",            "url": "https://investors.bio-rad.com/press-releases/default.aspx"},
    {"company": "Diasorin",                "url": "https://int.diasorin.com/en/investors/finance/press-releases"},
    {"company": "Dr. Fooke",               "url": "https://www.fooke-labs.de/veranstaltungen?lang=en"},
    {"company": "SVAR Lifesciences",       "url": "https://www.svarlifescience.com/news"},
    {"company": "Euroimmun",               "url": "https://www.euroimmun.us/recent-news"},
    {"company": "Eurospital",              "url": "https://www.eurospital.com/news/"},
    {"company": "HOB Biotech",             "url": "http://en.hob-biotech.com/media-55.html"},
    {"company": "Hycor Biomedical",        "url": "https://www.hycorbiomedical.com/news/categories/hycor-news-archive-2023"},
    {"company": "MBL",                     "url": "https://www.mblbio.com/e/ir/press.html#"},
    {"company": "JSR Corporation",         "url": "https://www.jsr.co.jp/jsr_e/news/2026/"},
    {"company": "QuidelOrtho",             "url": "https://ir.quidelortho.com/home/default.aspx"},
    {"company": "C-Luminary",              "url": "https://www.c-luminary.com/list-46-1.html"},
    {"company": "Nippon Chemiphar",        "url": "https://www.chemiphar.co.jp/english/"},
    {"company": "Human Diagnostics",       "url": "https://www.human.de/about-human/overview-human/news"},
    {"company": "Autobio Diagnostics",     "url": "https://en.autobio.com.cn/News/index/fid/3/cid/2/id/142.html"},
    {"company": "Siemens Healthineers",    "url": "https://www.siemens-healthineers.com/press/releases"},
    {"company": "Menarini",                "url": "https://www.menarini.com/en-us/news"},
    {"company": "Roche",                   "url": "https://www.roche.com/media/releases"},
    {"company": "Sebia",                   "url": "https://www.sebia.com/en-in/resources/"},
    {"company": "Revvity",                 "url": "https://news.revvity.com/press-announcements/press-releases/default.aspx"},
    {"company": "Zeus Scientific",         "url": "https://www.zeusscientific.com/news-events"},
    {"company": "Beckman (India)",         "url": "https://www.mybeckman.in/news"},
    {"company": "AliveDx",                 "url": "https://alivedx.com/news/"},
    {"company": "Werfen",                  "url": "https://www.werfen.com/en/news"},
    {"company": "Theradiag",               "url": "https://www.theradiag.com/en/category/press-releases/"},
    {"company": "Trinity Biotech",         "url": "https://www.trinitybiotech.com/category/press-releases/"},
    {"company": "BUHLMANN Laboratories",   "url": "https://www.buhlmannlabs.ch/news/"},
    {"company": "IDS plc",                 "url": "https://www.idsplc.com/blog/"},
    {"company": "Gold Standard Diagnostics","url": "https://www.goldstandarddiagnostics.com/news-event/category/news"},
    {"company": "R-Biopharm",              "url": "https://clinical.r-biopharm.com/news/"},
    {"company": "MADx",                    "url": "https://www.macroarraydx.com/news/overview"},
    {"company": "Aimmune",                 "url": "https://www.aimmune.com/news-releases"},
    {"company": "Stallergenes Greer",      "url": "https://www.stallergenesgreer.com/press-releases#"},
    {"company": "Allergy Therapeutics",    "url": "https://www.allergytherapeutics.com/investors/regulatory-news/"},
    {"company": "ALK-Abelló",             "url": "https://ir.alk.net/news-events/company-releases"},
    {"company": "DBV Technologies",        "url": "https://dbv-technologies.com/investor-overview/news/"},
    {"company": "Eurospital (EGL)",        "url": "https://egl.eurospital.com/en/news-and-events/"},
]


# ══════════════════════════════════════════════════════════════════════════════
# FILTERS
# ══════════════════════════════════════════════════════════════════════════════
def _contains_any(text, vocab):
    t = (text or "").lower()
    return any(k.lower() in t for k in vocab)

def keyword_is_relevant(text):
    if not text:
        return False
    if _contains_any(text, EXCLUDE_KEYWORDS):
        return False
    return _contains_any(text, RELEVANT_KEYWORDS)

def classify_news_type(text):
    t = (text or "").lower()
    for cat, triggers in NEWS_TYPE_RULES.items():
        if any(trig.lower() in t for trig in triggers):
            return cat
    return "Other"

def classify_hot(text):
    return "HOT" if _contains_any(text, HOT_KEYWORDS) else "Non-Hot"

def deduplicate(records):
    seen, out = set(), []
    for r in records:
        entity   = (r.get("screened_entity") or r.get("company") or "").lower().strip()
        category = (r.get("screening_category") or "").lower().strip()
        headline = (r.get("headline") or "").lower().strip()[:80]
        key = (category, entity, headline)
        if key in seen:
            continue
        seen.add(key)
        out.append(r)
    return out


# ══════════════════════════════════════════════════════════════════════════════
# GOOGLE NEWS
# ══════════════════════════════════════════════════════════════════════════════
_GNEWS = "https://news.google.com/rss/search?q={q}&hl=en-US&gl=US&ceid=US:en"

def _parse_date(entry):
    for f in ("published_parsed","updated_parsed"):
        t = entry.get(f)
        if t:
            try:
                return datetime(*t[:6]).date()
            except Exception:
                pass
    for f in ("published","updated"):
        raw = entry.get(f)
        if raw:
            try:
                return dateparser.parse(raw).date()
            except Exception:
                pass
    return None

def _fetch_gnews(entities, start_date, end_date, label, key_fn):
    records = []
    total = len(entities)
    for idx, entity in enumerate(entities, 1):
        name = entity["name"] if isinstance(entity, dict) else entity
        owner = entity.get("company", name) if isinstance(entity, dict) else entity
        _log(f"  [{label}] {idx}/{total} — {name}")
        url = _GNEWS.format(q=quote_plus(f'"{name}" diagnostics OR allergy OR immunology'))
        try:
            feed = feedparser.parse(url)
        except Exception as e:
            logger.warning(f"Feed parse failed {name}: {e}")
            continue
        for entry in feed.entries:
            pub = _parse_date(entry)
            if pub is None or pub < start_date or pub > end_date:
                continue
            headline = (entry.get("title") or "").strip()
            link     = (entry.get("link")  or "").strip()
            if not headline or not link:
                continue
            records.append({
                "company": owner, "screened_entity": name,
                "screening_category": label,
                "headline": headline,
                "date": pub.strftime("%Y-%m-%d"),
                "url": link, "source_type": "Google News",
            })
    return records

def fetch_google_news(start_date, end_date):
    _log(f"\n── Competitor Screening ({len(COMPANIES)} companies) ──")
    return _fetch_gnews(COMPANIES, start_date, end_date, "Competitor Screening", lambda x: x)

def fetch_product_news(start_date, end_date):
    _log(f"\n── Product Screening ({len(PRODUCTS)} products) ──")
    return _fetch_gnews(PRODUCTS, start_date, end_date, "Product Screening", lambda x: x)


# ══════════════════════════════════════════════════════════════════════════════
# PR SCRAPER
# ══════════════════════════════════════════════════════════════════════════════
_DATE_PATS = [
    re.compile(r"\b(\d{4}-\d{2}-\d{2})\b"),
    re.compile(r"\b((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{1,2},?\s+\d{4})\b", re.I),
    re.compile(r"\b(\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4})\b", re.I),
    re.compile(r"\b(\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4})\b"),
]
_DATE_HINTS = ("date","published","pub-date","pubdate","timestamp","time","posted")
_NOISE_RE = re.compile(
    r"^(home|news|press|releases?|menu|search|login|sign in|sign up|subscribe|"
    r"contact|cookie|privacy|terms|imprint|read more|learn more|more|back|next|"
    r"previous|all|view all|see all|share|download|english|deutsch|français|"
    r"about|careers|investors|products|solutions|support|events)$", re.I)

def _try_parse(raw):
    if not raw:
        return None
    raw = raw.strip()
    if len(raw) < 6 or len(raw) > 60:
        return None
    try:
        d = dateparser.parse(raw, fuzzy=True).date()
    except Exception:
        return None
    if d.year < 2000 or d.year > date.today().year + 1:
        return None
    return d

def _date_from_attrs(el):
    if not hasattr(el, "get"):
        return None
    for attr in ("datetime","data-date","data-published","data-time","content"):
        v = el.get(attr)
        if v:
            d = _try_parse(v)
            if d:
                return d
    return None

def _date_from_text(text):
    for pat in _DATE_PATS:
        m = pat.search(text or "")
        if m:
            d = _try_parse(m.group(1))
            if d:
                return d
    return None

def _date_from_element(el, max_text=1500):
    if el is None or not hasattr(el, "find_all"):
        return None
    d = _date_from_attrs(el)
    if d:
        return d
    for t in el.find_all("time", limit=6):
        d = _date_from_attrs(t) or _try_parse(t.get_text(" ", strip=True))
        if d:
            return d
    for cand in el.find_all(True, limit=40):
        cls = " ".join(cand.get("class") or []).lower()
        if not any(h in cls for h in _DATE_HINTS):
            continue
        d = _date_from_attrs(cand) or _try_parse(cand.get_text(" ", strip=True))
        if d:
            return d
    return _date_from_text(el.get_text(" ", strip=True)[:max_text])

def _find_date_for_anchor(a):
    d = _date_from_element(a, max_text=400)
    if d:
        return d
    parent = a.parent
    for _ in range(4):
        if parent is None or parent.name in ("body","html"):
            break
        d = _date_from_element(parent, max_text=1500)
        if d:
            return d
        parent = parent.parent
    for fn in (a.find_previous_sibling, a.find_next_sibling):
        sib = fn()
        for _ in range(3):
            if sib is None:
                break
            d = _date_from_element(sib, max_text=400)
            if d:
                return d
            sib = sib.find_previous_sibling() if fn == a.find_previous_sibling else sib.find_next_sibling()
    return None

def _is_internal(href, base_netloc):
    if not href or href.startswith(("javascript:","mailto:","tel:","#")):
        return False
    if href.startswith("/"):
        return True
    p = urlparse(href)
    if not p.netloc:
        return True
    return base_netloc.split(".")[-2:] == p.netloc.split(".")[-2:]

def _collect_candidates(soup, base_url):
    base_netloc = urlparse(base_url).netloc
    out, seen = [], set()
    for a in soup.find_all("a", href=True, limit=2000):
        href = a["href"].strip()
        if not _is_internal(href, base_netloc):
            continue
        headline = None
        for tag in ("h1","h2","h3","h4","h5"):
            h = a.find(tag)
            if h:
                txt = h.get_text(" ", strip=True)
                if 15 <= len(txt) <= 500:
                    headline = txt
                    break
        if not headline:
            txt = a.get_text(" ", strip=True)
            if 15 <= len(txt) <= 500 and not _NOISE_RE.match(txt):
                headline = txt
        if not headline:
            continue
        key = headline.lower()[:80]
        if key in seen:
            continue
        seen.add(key)
        out.append((a, headline, urljoin(base_url, href)))
    return out

def _classify_failure(exc):
    if isinstance(exc, requests.HTTPError) and exc.response is not None:
        code = exc.response.status_code
        if code == 404: return "HTTP 404 — Not found"
        if code == 403: return "HTTP 403 — Forbidden"
        if 500 <= code < 600: return f"HTTP {code} — Server error"
        return f"HTTP {code}"
    if isinstance(exc, requests.ConnectionError):
        s = str(exc).lower()
        if any(k in s for k in ("nodename","getaddrinfo")): return "DNS failure"
        if "ssl" in s or "certificate" in s: return "SSL error"
        return "Connection error"
    if isinstance(exc, requests.Timeout): return "Timeout"
    if isinstance(exc, requests.TooManyRedirects): return "Too many redirects"
    return type(exc).__name__

def scrape_pr_websites(start_date, end_date, include_undated=True, undated_cap=15):
    _log(f"\n── PR Website Scraping ({len(PR_WEBSITES)} sites) ──")
    records, failures = [], []
    today_str = date.today().strftime("%Y-%m-%d")

    for idx, site in enumerate(PR_WEBSITES, 1):
        company, url = site["company"], site["url"]
        _log(f"  [PR] {idx}/{len(PR_WEBSITES)} — {company}")
        if "linkedin.com" in urlparse(url).netloc.lower():
            failures.append({"company": company, "url": url, "category": "Blocked (LinkedIn)", "error": "LinkedIn blocks scraping"})
            continue
        resp, last_exc = None, None
        for attempt in range(2):
            try:
                resp = requests.get(url, headers=HEADERS, timeout=REQUEST_TIMEOUT)
                resp.raise_for_status()
                last_exc = None
                break
            except Exception as e:
                last_exc = e
                if attempt == 0:
                    time.sleep(1.5)
        if last_exc or resp is None:
            cat = _classify_failure(last_exc or Exception("Unknown"))
            failures.append({"company": company, "url": url, "category": cat, "error": str(last_exc)[:200]})
            continue
        try:
            soup = BeautifulSoup(resp.text, "lxml")
        except Exception as e:
            failures.append({"company": company, "url": url, "category": "Parse error", "error": str(e)[:200]})
            continue
        candidates = _collect_candidates(soup, url)
        in_range = 0
        dated_count = 0
        undated_pool = []
        for a, headline, link in candidates:
            d = _find_date_for_anchor(a)
            if d is None:
                undated_pool.append((headline, link))
                continue
            dated_count += 1
            if d < start_date or d > end_date:
                continue
            records.append({
                "company": company, "screened_entity": company,
                "screening_category": "PR Site Screening",
                "headline": headline, "date": d.strftime("%Y-%m-%d"),
                "url": link, "source_type": "PR Website", "date_estimated": False,
            })
            in_range += 1
        if include_undated and dated_count == 0:
            for headline, link in undated_pool[:undated_cap]:
                records.append({
                    "company": company, "screened_entity": company,
                    "screening_category": "PR Site Screening",
                    "headline": headline, "date": today_str,
                    "url": link, "source_type": "PR Website", "date_estimated": True,
                })
        _log(f"    → candidates={len(candidates)} dated={dated_count} in_range={in_range} undated_used={min(len(undated_pool), undated_cap) if include_undated and dated_count==0 else 0}")
    return records, failures


# ══════════════════════════════════════════════════════════════════════════════
# AI FILTER (Groq)
# ══════════════════════════════════════════════════════════════════════════════
GROQ_ENDPOINT = "https://api.groq.com/openai/v1/chat/completions"
DEFAULT_MODEL  = "llama-3.3-70b-versatile"

AI_SYSTEM_PROMPT = """You are an intelligence analyst preparing a Daily Newsletter (DNL) for a specialty diagnostics business unit focused on immunodiagnostics, allergy, and autoimmune testing.

Your job is to classify headlines as RELEVANT or NOT RELEVANT for the Daily Newsletter (DNL).

IN SCOPE (relevant):
- Immunodiagnostics, allergy diagnostics, autoimmune diagnostics
- Immunoassays (ELISA, lateral flow, multiplex, serology)
- Diagnostic platforms / analyzers / reagents / test kits
- Infectious disease diagnostics, point-of-care (POC) diagnostics, IVD
- Laboratory solutions
- Allergen immunotherapy / allergy therapeutics (ALK, Stallergenes Greer, Aimmune, DBV, Allergy Therapeutics are in scope)
- Strategic events on monitored companies: M&A, partnerships, regulatory approvals (FDA, CE, 510k), product launches, conference participation, leadership changes, strategic financial results

OUT OF SCOPE (not relevant):
- Generic pharma news (cardiovascular, oncology, diabetes, obesity)
- Biosimilars or unrelated drug approvals
- General hospital / healthcare system news
- Unrelated clinical trials for drugs
- Mental health / psychiatry / neurology / orthopedic / dermatology / ophthalmology drugs
- Generic medical news with no link to diagnostics or allergy

Return ONLY a JSON array, no prose, no markdown fences. Each element:
{"headline": "<exact headline you were given>", "relevant": true|false, "reason": "<one short sentence>"}
"""

def _extract_json_array(text):
    text = re.sub(r"^```(?:json)?", "", text.strip(), flags=re.I)
    text = re.sub(r"```$", "", text.strip())
    s, e = text.find("["), text.rfind("]")
    if s == -1 or e == -1 or e <= s:
        raise ValueError("no JSON array found")
    return json.loads(text[s:e+1])

def _parse_verdicts(text):
    try:
        obj = json.loads(text.strip())
        if isinstance(obj, list):
            return obj
        if isinstance(obj, dict):
            for k in ("results","verdicts","classifications","data"):
                if k in obj and isinstance(obj[k], list):
                    return obj[k]
            for v in obj.values():
                if isinstance(v, list):
                    return v
    except Exception:
        pass
    return _extract_json_array(text)

def ai_filter_batch(records, failures, start_date, end_date, batch_size=20):
    api_key = os.getenv("GROQ_API_KEY")
    model   = os.getenv("GROQ_MODEL", DEFAULT_MODEL)

    # Build augmented system prompt
    prompt = AI_SYSTEM_PROMPT + f"\n\nDATE RANGE: {start_date} → {end_date}."
    if failures:
        lines = "\n".join(f"- {f['company']}: {f['url']}" for f in failures[:60])
        prompt += (
            "\n\nUNREACHABLE PR SITES (could not be scraped):\n" + lines +
            "\n\nFor these companies, if you know of significant announcements within the "
            "date range, include a synthetic entry with ai_synthesized: true. Never fabricate."
        )

    if not api_key:
        _log("[AI] GROQ_API_KEY not set — using keyword-only filter")
        for r in records:
            r["ai_relevant"] = True
            r["ai_reason"]   = "AI skipped (no key)"
        return records

    _log(f"\n── AI Filter ({model}) — {len(records)} records in {-(-len(records)//batch_size)} batches ──")
    kept = []
    total_batches = -(-len(records) // batch_size)

    for b_idx, start in enumerate(range(0, len(records), batch_size), 1):
        batch   = records[start:start+batch_size]
        payload = [{"company": r.get("company",""), "headline": r.get("headline","")} for r in batch]
        user_msg = (
            f"Classify these {len(payload)} headlines. "
            'Return {"results":[{"headline":"...","relevant":true|false,"reason":"..."}]}.\n\n'
            + json.dumps(payload, ensure_ascii=False)
        )
        _log(f"  [AI] batch {b_idx}/{total_batches} ({len(batch)} headlines) …")
        try:
            r = requests.post(
                GROQ_ENDPOINT,
                headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                json={
                    "model": model, "temperature": 0.1, "max_tokens": 4096,
                    "response_format": {"type": "json_object"},
                    "messages": [
                        {"role": "system", "content": prompt + '\n\nReturn the JSON array under a key called "results".'},
                        {"role": "user",   "content": user_msg},
                    ],
                },
                timeout=90,
            )
            r.raise_for_status()
            verdicts = _parse_verdicts(r.json()["choices"][0]["message"]["content"])
        except Exception as e:
            _log(f"  [AI] batch {b_idx} FAILED ({e}) — keeping all (fail-open)")
            for rec in batch:
                rec["ai_relevant"] = True
                rec["ai_reason"]   = f"AI error: {type(e).__name__}"
                kept.append(rec)
            continue

        by_headline   = {v["headline"].strip().lower(): v for v in verdicts if isinstance(v, dict) and "headline" in v and not v.get("ai_synthesized")}
        synthesized   = [v for v in verdicts if isinstance(v, dict) and v.get("ai_synthesized")]

        for rec in batch:
            v = by_headline.get((rec.get("headline") or "").strip().lower())
            if v is None or bool(v.get("relevant", True)):
                rec["ai_relevant"] = True
                rec["ai_reason"]   = str((v or {}).get("reason","no verdict"))[:200]
                kept.append(rec)
            else:
                rec["ai_relevant"] = False
                rec["ai_reason"]   = str(v.get("reason",""))[:200]

        today_str = date.today().strftime("%Y-%m-%d")
        for v in synthesized:
            kept.append({
                "company": str(v.get("company","Unknown"))[:120],
                "screened_entity": str(v.get("company","Unknown"))[:120],
                "screening_category": "PR Site Screening",
                "headline": str(v.get("headline",""))[:500],
                "date": today_str, "url": "",
                "source_type": "AI Knowledge (site unreachable)",
                "ai_relevant": True, "ai_reason": "AI-synthesized",
                "ai_synthesized": True,
            })
        _log(f"  [AI] batch {b_idx} done — kept {sum(1 for r in batch if r.get('ai_relevant',True))}/{len(batch)}, synthesized {len(synthesized)}")

    _log(f"[AI] Total kept: {len(kept)} of {len(records)}")
    return kept


# ══════════════════════════════════════════════════════════════════════════════
# EXCEL EXPORTER
# ══════════════════════════════════════════════════════════════════════════════
OUTPUT_DIR = "output"
COLUMNS = [
    ("Serial Number",      "serial_number",      8),
    ("Company Name",       "company",            24),
    ("Screened Entity",    "screened_entity",    24),
    ("Screening Category", "screening_category", 22),
    ("Date",               "date",               13),
    ("News Type",          "news_type",          18),
    ("Headline",           "headline",           60),
    ("Source Link",        "url",                40),
    ("Source Type",        "source_type",        14),
    ("Hot vs Non-Hot",     "hot",                14),
    ("Date Collected",     "date_collected",     14),
]
NAVY       = "1B3A5C"
HOT_YELLOW = "FFF2CC"
ALT_GREY   = "F5F7FA"
_THIN      = Side(style="thin", color="D0D7DE")
_BORDER    = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

def export_to_excel(records, start_date, end_date):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"DNL_{start_date}_{end_date}_{ts}.xlsx"
    filepath = os.path.abspath(os.path.join(OUTPUT_DIR, filename))

    wb = Workbook()
    ws = wb.active
    ws.title = "DNL Report"

    # Header row
    hf = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor=NAVY)
    ha = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ci, (label, _, width) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=ci, value=label)
        c.font, c.fill, c.alignment, c.border = hf, hfill, ha, _BORDER
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    # Data rows
    for i, rec in enumerate(records):
        is_hot = rec.get("hot") == "HOT"
        alt    = i % 2 == 1
        fill   = PatternFill("solid", fgColor=HOT_YELLOW) if is_hot else (PatternFill("solid", fgColor=ALT_GREY) if alt else None)
        ra     = Alignment(vertical="center", wrap_text=True)
        for ci, (_, key, _) in enumerate(COLUMNS, 1):
            val  = rec.get(key, "") or (rec.get("company","") if key == "screened_entity" else "")
            cell = ws.cell(row=i+2, column=ci, value=val)
            cell.alignment, cell.border = ra, _BORDER
            if fill:
                cell.fill = fill
            if key == "url" and val:
                cell.hyperlink = val
                cell.font = Font(name="Calibri", size=10, color="0563C1", underline="single")
            else:
                cell.font = Font(name="Calibri", size=10, bold=(is_hot and key=="hot"))
        ws.row_dimensions[i+2].height = 45

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(records)+1)}"

    # Summary sheet
    ss = wb.create_sheet("Summary")
    ss["A1"] = "DNL Report — Summary"
    ss["A1"].font  = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ss["A1"].fill  = PatternFill("solid", fgColor=NAVY)
    ss["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ss.merge_cells("A1:B1")
    ss.row_dimensions[1].height = 28
    bf = Font(name="Calibri", size=11, bold=True)
    ss["A3"], ss["B3"] = "Date Range",     f"{start_date} → {end_date}"
    ss["A4"], ss["B4"] = "Generated",      datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ss["A5"], ss["B5"] = "Total Records",  len(records)
    ss["A6"], ss["B6"] = "HOT Items",      sum(1 for r in records if r.get("hot")=="HOT")
    ss["A7"], ss["B7"] = "Product Results",sum(1 for r in records if r.get("screening_category")=="Product Screening")
    for row in range(3, 8):
        ss.cell(row=row, column=1).font = bf
    ss["A8"], ss["B8"] = "News Type", "Count"
    ss["A8"].font = ss["B8"].font = bf
    counts = {}
    for r in records:
        nt = r.get("news_type","Other")
        counts[nt] = counts.get(nt,0) + 1
    for ri, (nt, c) in enumerate(sorted(counts.items(), key=lambda x:-x[1]), 9):
        ss.cell(row=ri, column=1, value=nt)
        ss.cell(row=ri, column=2, value=c)
    ss.column_dimensions["A"].width = 28
    ss.column_dimensions["B"].width = 16

    wb.save(filepath)
    return filepath


# ══════════════════════════════════════════════════════════════════════════════
# MAIN PIPELINE
# ══════════════════════════════════════════════════════════════════════════════
def run(start_date, end_date, use_ai=True, use_google=True, use_pr=True, use_products=True, include_undated=True):
    _log(f"\n{'='*60}")
    _log(f"DNL Automation — {start_date} → {end_date}")
    _log(f"{'='*60}")
    _log(f"Sources: Google={use_google} PR={use_pr} Products={use_products} AI={use_ai}")

    raw_google   = fetch_google_news(start_date, end_date)   if use_google   else []
    raw_pr, fails= scrape_pr_websites(start_date, end_date, include_undated=include_undated) if use_pr else ([], [])
    raw_products = fetch_product_news(start_date, end_date)  if use_products else []

    all_records  = raw_google + raw_pr + raw_products
    _log(f"\n── Keyword filter ──")
    kw_filtered  = [r for r in all_records if keyword_is_relevant(r.get("headline",""))]
    _log(f"  Kept {len(kw_filtered)} of {len(all_records)}")

    if use_ai:
        ai_filtered = ai_filter_batch(kw_filtered, fails, start_date, end_date)
    else:
        ai_filtered = kw_filtered
        for r in ai_filtered:
            r.setdefault("ai_relevant", True)
            r.setdefault("ai_reason", "AI disabled")

    _log(f"\n── Deduplicate ──")
    deduped = deduplicate(ai_filtered)
    _log(f"  {len(deduped)} unique ({len(ai_filtered)-len(deduped)} removed)")

    _log(f"\n── Classify ──")
    today_str = datetime.now().strftime("%Y-%m-%d")
    hot_count = 0
    for i, r in enumerate(deduped, 1):
        r["news_type"]     = classify_news_type(r.get("headline",""))
        r["hot"]           = classify_hot(r.get("headline",""))
        r["serial_number"] = i
        r["date_collected"]= today_str
        if r["hot"] == "HOT":
            hot_count += 1
    _log(f"  {len(deduped)} records classified — {hot_count} HOT")

    # Export Excel
    _log(f"\n── Export ──")
    excel_path = export_to_excel(deduped, start_date, end_date)
    _log(f"  Excel: {excel_path}")

    # Export failures CSV
    fails_path = None
    if fails:
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        fails_path = os.path.join(OUTPUT_DIR, f"DNL_{start_date}_{end_date}_{ts}_failures.csv")
        pd.DataFrame(fails).to_csv(fails_path, index=False)
        _log(f"  Failures CSV: {fails_path} ({len(fails)} sites failed)")

    # Write run_summary.md (picked up by GitHub Actions job summary)
    summary_lines = [
        f"## DNL Run Summary",
        f"**Date range:** {start_date} → {end_date}  ",
        f"**Generated:** {today_str}  ",
        f"",
        f"| Metric | Value |",
        f"|--------|-------|",
        f"| Total records | **{len(deduped)}** |",
        f"| HOT items | **{hot_count}** |",
        f"| Google News records | {sum(1 for r in deduped if r.get('source_type')=='Google News')} |",
        f"| PR Website records | {sum(1 for r in deduped if r.get('source_type')=='PR Website')} |",
        f"| Product Screening records | {sum(1 for r in deduped if r.get('screening_category')=='Product Screening')} |",
        f"| PR sites failed | {len(fails)} |",
        f"| AI filter used | {'Yes ✅' if use_ai and os.getenv('GROQ_API_KEY') else 'No (keyword-only)'} |",
        f"",
        f"### News Type Breakdown",
        f"| Type | Count |",
        f"|------|-------|",
    ]
    counts = {}
    for r in deduped:
        nt = r.get("news_type","Other")
        counts[nt] = counts.get(nt,0) + 1
    for nt, c in sorted(counts.items(), key=lambda x:-x[1]):
        summary_lines.append(f"| {nt} | {c} |")

    summary_lines += [
        f"",
        f"### HOT Items",
    ]
    hot_items = [r for r in deduped if r.get("hot")=="HOT"]
    if hot_items:
        for r in hot_items:
            summary_lines.append(f"- **{r.get('company','')}** — {r.get('headline','')} _{r.get('news_type','')}_")
    else:
        summary_lines.append("_No HOT items this run._")

    summary_path = os.path.join(OUTPUT_DIR, "run_summary.md")
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    Path(summary_path).write_text("\n".join(summary_lines), encoding="utf-8")
    _log(f"  Summary: {summary_path}")

    # Also write to GITHUB_STEP_SUMMARY if available
    gss = os.environ.get("GITHUB_STEP_SUMMARY")
    if gss:
        try:
            with open(gss, "w", encoding="utf-8") as f:
                f.write("\n".join(summary_lines))
        except Exception:
            pass

    _log(f"\n{'='*60}")
    _log(f"DONE — {len(deduped)} records, {hot_count} HOT")
    _log(f"{'='*60}\n")

    return deduped, fails, excel_path


# ══════════════════════════════════════════════════════════════════════════════
# CLI ENTRY POINT
# ══════════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="DNL Automation — GitHub Actions runner")
    parser.add_argument("--start",       type=str,  help="Start date YYYY-MM-DD (default: 3 days ago)")
    parser.add_argument("--end",         type=str,  help="End date YYYY-MM-DD (default: today)")
    parser.add_argument("--days",        type=int,  default=3, help="Lookback days (used if --start not set)")
    parser.add_argument("--no-ai",       action="store_true", help="Skip AI filter")
    parser.add_argument("--no-google",   action="store_true", help="Skip Google News")
    parser.add_argument("--no-pr",       action="store_true", help="Skip PR scraping")
    parser.add_argument("--no-products", action="store_true", help="Skip product screening")
    parser.add_argument("--no-undated",  action="store_true", help="Drop PR articles without a detectable date")
    args = parser.parse_args()

    end_date   = date.fromisoformat(args.end)   if args.end   else date.today()
    start_date = date.fromisoformat(args.start) if args.start else end_date - timedelta(days=args.days)

    records, fails, excel = run(
        start_date      = start_date,
        end_date        = end_date,
        use_ai          = not args.no_ai,
        use_google      = not args.no_google,
        use_pr          = not args.no_pr,
        use_products    = not args.no_products,
        include_undated = not args.no_undated,
    )

    sys.exit(0 if records is not None else 1)
